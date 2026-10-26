import os
import tempfile

os.environ["SIGT_DATA_DIR"] = tempfile.mkdtemp()  # SQLite si no hay SIGT_DATABASE_URL
os.environ["SIGT_ADMIN_USER"] = "admin"
os.environ["SIGT_ADMIN_PASSWORD"] = "clave-admin-123"

import pytest
from fastapi.testclient import TestClient

from app import calc
from app.main import app

CLAVE = "clave-segura-123"


def test_semaforo_directo_e_inverso():
    assert calc.semaforo(95, 90, 80, "mayor") == "verde"
    assert calc.semaforo(85, 90, 80, "mayor") == "amarillo"
    assert calc.semaforo(70, 90, 80, "mayor") == "rojo"
    assert calc.semaforo(2, 3, 5, "menor") == "verde"
    assert calc.semaforo(4, 3, 5, "menor") == "amarillo"
    assert calc.semaforo(6, 3, 5, "menor") == "rojo"


def test_formula_segura():
    assert calc.evaluar("a / b * 100", {"a": 1, "b": 4}) == 25
    with pytest.raises(calc.FormulaError):
        calc.evaluar("a / b", {"a": 1, "b": 0})
    with pytest.raises(calc.FormulaError):
        calc.validar_formula("__import__('os').system('ls')", ["a"])


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def entrar(c, user, clave=CLAVE):
    c.cookies.clear()
    r = c.post("/api/login", json={"username": user, "password": clave})
    assert r.status_code == 200, r.text


@pytest.fixture(scope="module")
def datos(client):
    """Base vacía → el administrador crea área, usuarios e indicador."""
    c = client
    entrar(c, "admin", "clave-admin-123")
    assert c.get("/api/indicadores").json() == []
    area = c.post("/api/areas", json={"nombre": "Sistemas"}).json()["id"]
    otra = c.post("/api/areas", json={"nombre": "Producción"}).json()["id"]
    for u, rol, a in [("jefe.sis", "jefe", area), ("jefe.pro", "jefe", otra), ("gerente", "gerente", None)]:
        assert c.post("/api/usuarios", json={"username": u, "nombre": f"Usuario {u}", "rol": rol,
                                             "area_id": a, "password": CLAVE}).status_code == 201
    ind = c.post("/api/indicadores", json={
        "codigo": "sis-03", "nombre": "Tiempo de resolución", "area_id": area, "unidad": "días",
        "frecuencia": "mensual", "sentido": "menor", "formula": "dias / tickets", "meta": 2, "limite": 4,
        "variables": [{"clave": "dias", "etiqueta": "Días acumulados"}, {"clave": "tickets", "etiqueta": "Tickets"}]})
    assert ind.status_code == 201
    return {"area": area, "ind": ind.json()["id"]}


def test_sin_datos_demo_ni_endpoint_demo(client):
    assert client.get("/api/demo-usuarios").status_code in (404, 405)


def test_flujo_completo(client, datos):
    c, ind = client, datos["ind"]
    entrar(c, "jefe.sis")
    per = c.get(f"/api/indicadores/{ind}/periodos").json()["periodos"][0]["periodo"]
    body = {"indicador_id": ind, "periodo": per, "valores": {"dias": 300, "tickets": 50}}
    assert c.post("/api/mediciones", json=body).status_code == 422  # fuera de meta sin análisis
    body.update(analisis="Caída del enlace principal", plan_mejora="Enlace de respaldo y turno extra")
    r = c.post("/api/mediciones", json=body)
    assert r.status_code == 201 and r.json()["semaforo"] == "rojo"
    assert c.post("/api/mediciones", json=body).status_code == 409
    entrar(c, "jefe.pro")
    assert c.post("/api/calcular", json={"indicador_id": ind, "valores": {}}).status_code == 403
    assert c.get("/api/dashboard").status_code == 403
    assert c.get("/api/usuarios").status_code == 403
    entrar(c, "gerente")
    mid = r.json()["id"]
    assert c.post(f"/api/mediciones/{mid}/revision", json={"accion": "observar"}).status_code == 422
    assert c.post(f"/api/mediciones/{mid}/revision", json={"accion": "observar", "nota": "Falta fecha"}).status_code == 200
    entrar(c, "jefe.sis")
    assert c.post("/api/mediciones", json=body).json()["version"] == 2
    entrar(c, "gerente")
    assert c.post(f"/api/mediciones/{mid}/revision", json={"accion": "aprobar"}).json()["estado"] == "aprobado"
    assert c.post(f"/api/mediciones/{mid}/notas", json={"texto": "Verificar eficacia"}).status_code == 201
    assert c.get("/api/bitacora/verificar").status_code == 403
    entrar(c, "admin", "clave-admin-123")
    assert c.get("/api/bitacora/verificar").json()["ok"] is True


def test_clave_invalida_sesiones_y_limite_de_intentos(client, datos):
    c = client
    entrar(c, "jefe.sis")
    assert c.post("/api/cambiar-clave", json={"actual": "incorrecta", "nueva": "otra-clave-123"}).status_code == 422
    antes = c.cookies.get("sigt_sesion")
    assert c.post("/api/cambiar-clave", json={"actual": CLAVE, "nueva": "otra-clave-123"}).status_code == 200
    assert c.get("/api/me").status_code == 200  # la sesión actual se renueva
    c.cookies.clear()
    c.cookies.set("sigt_sesion", antes)
    assert c.get("/api/me").status_code == 401  # la sesión anterior queda invalidada
    for _ in range(5):
        c.post("/api/login", json={"username": "jefe.pro", "password": "mala"})
    assert c.post("/api/login", json={"username": "jefe.pro", "password": CLAVE}).status_code == 429
    from app.main import limite_login
    limite_login._fallos.clear()


def test_admin_integridad(client, datos):
    c = client
    entrar(c, "admin", "clave-admin-123")
    assert c.delete(f"/api/areas/{datos['area']}").status_code == 409  # tiene indicadores
    assert c.delete(f"/api/indicadores/{datos['ind']}").json()["archivado"] is True
    assert c.post("/api/usuarios", json={"username": "x.jefe", "nombre": "Sin Area", "rol": "jefe",
                                         "password": CLAVE}).status_code == 422
    me = c.get("/api/me").json()
    assert c.put(f"/api/usuarios/{me['id']}", json={"nombre": "Admin SGC", "rol": "admin", "activo": False}).status_code == 422
    from app.database import SessionLocal
    from app.models import Bitacora
    with SessionLocal() as db:
        b = db.get(Bitacora, 1)
        b.accion = "X"
        with pytest.raises(RuntimeError):
            db.flush()


def test_admin_es_superusuario_con_segregacion_de_funciones(client):
    c = client
    entrar(c, "admin", "clave-admin-123")
    area = c.post("/api/areas", json={"nombre": "Calidad"}).json()["id"]
    ind = c.post("/api/indicadores", json={
        "codigo": "CAL-01", "nombre": "Lotes conformes", "area_id": area, "unidad": "%",
        "frecuencia": "mensual", "sentido": "mayor", "formula": "ok / total * 100", "meta": 98, "limite": 95,
        "variables": [{"clave": "ok", "etiqueta": "Conformes"}, {"clave": "total", "etiqueta": "Analizados"}]}).json()["id"]
    pr = c.post("/api/formula/probar", json={"formula": "ok / total * 100", "valores": {"ok": 99, "total": 100},
                                             "meta": 98, "limite": 95, "sentido": "mayor"}).json()
    assert pr["ok"] and pr["semaforo"] == "verde"
    assert c.post("/api/formula/validar", json={"formula": "ok * 2", "claves": ["ok", "total"]}).json()["sin_usar"] == ["total"]
    per = c.get(f"/api/indicadores/{ind}/periodos").json()["periodos"][1]["periodo"]
    r = c.post("/api/mediciones", json={"indicador_id": ind, "periodo": per, "valores": {"ok": 99, "total": 100}})
    assert r.status_code == 201  # el admin puede cargar en cualquier área
    mid = r.json()["id"]
    assert c.post(f"/api/mediciones/{mid}/revision", json={"accion": "aprobar"}).status_code == 403  # carga propia
    entrar(c, "gerente")
    assert c.post(f"/api/mediciones/{mid}/revision", json={"accion": "aprobar"}).status_code == 200
    # con cargas, la ficha no admite cambiar frecuencia ni claves
    entrar(c, "admin", "clave-admin-123")
    ficha = c.get("/api/indicadores").json()
    ficha = next(i for i in ficha if i["id"] == ind)
    assert ficha["n_mediciones"] == 1
    base = {k: ficha[k] for k in ("codigo", "nombre", "objetivo", "area_id", "unidad", "sentido", "formula",
                                  "variables", "meta", "limite", "activo")}
    assert c.put(f"/api/indicadores/{ind}", json={**base, "frecuencia": "trimestral"}).status_code == 422
    malas = [{"clave": "ok", "etiqueta": "Conformes"}, {"clave": "tot", "etiqueta": "Analizados"}]
    assert c.put(f"/api/indicadores/{ind}", json={**base, "frecuencia": "mensual", "variables": malas,
                                                  "formula": "ok / tot * 100"}).status_code == 422
    assert c.put(f"/api/indicadores/{ind}", json={**base, "frecuencia": "mensual", "meta": 97}).status_code == 200


def test_formulas_complejas_constantes_y_opcionales(client):
    c = client
    entrar(c, "admin", "clave-admin-123")
    area = c.get("/api/areas").json()[0]["id"]
    ficha = {"codigo": "ssa-01", "nombre": "Índice de frecuencia", "area_id": area, "unidad": "índice",
             "frecuencia": "anual", "sentido": "menor", "decimales": 3, "meta": 5, "limite": 10,
             "formula": "si(horas > 0, accidentes × factor ÷ horas, 0) + extra",
             "variables": [{"clave": "accidentes", "etiqueta": "Accidentes"}, {"clave": "horas", "etiqueta": "Horas-hombre"},
                           {"clave": "extra", "etiqueta": "Ajuste", "opcional": True, "defecto": 1, "negativo": True}],
             "constantes": [{"clave": "factor", "etiqueta": "Factor por millón", "valor": 1000000}]}
    assert c.post("/api/indicadores", json={**ficha, "formula": "accidentes / 0 + noexiste"}).status_code == 422
    assert c.post("/api/indicadores", json={**ficha, "variables": [{"clave": "si", "etiqueta": "Reservada"}],
                                            "formula": "si"}).status_code == 422
    r = c.post("/api/indicadores", json=ficha)
    assert r.status_code == 201, r.text
    ind = r.json()
    assert ind["frecuencia"] == "anual" and ind["decimales"] == 3
    entrar(c, "admin", "clave-admin-123")
    per = c.get(f"/api/indicadores/{ind['id']}/periodos").json()["periodos"][1]["periodo"]
    assert len(per) == 4  # período anual
    # horas = 0 no rompe (rama protegida) y la variable opcional toma su valor por defecto
    calc_ = c.post("/api/calcular", json={"indicador_id": ind["id"], "valores": {"accidentes": 2, "horas": 0, "extra": None}}).json()
    assert calc_["ok"] and calc_["resultado"] == 1.0
    calc_ = c.post("/api/calcular", json={"indicador_id": ind["id"], "valores": {"accidentes": 2, "horas": 400000, "extra": -0.5}}).json()
    assert calc_["resultado"] == 4.5 and calc_["semaforo"] == "verde"
    r = c.post("/api/mediciones", json={"indicador_id": ind["id"], "periodo": per,
                                       "valores": {"accidentes": 2, "horas": 400000, "extra": -0.5}})
    assert r.status_code == 201 and r.json()["formula"].startswith("si(") and r.json()["constantes"][0]["valor"] == 1000000


def test_reportes_excel_y_bitacora_solo_admin(client):
    from io import BytesIO
    from openpyxl import load_workbook
    c = client
    entrar(c, "admin", "clave-admin-123")
    r = c.get("/api/reportes/general.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Resumen", "Mediciones", "Tendencias", "Plan de mejora", "Fichas técnicas"]
    assert wb["Mediciones"].max_row > 4 and wb["Fichas técnicas"].max_row > 4
    assert c.get("/api/reportes/bitacora.xlsx").status_code == 200
    assert c.get("/api/bitacora").status_code == 200
    for u in ("gerente", "jefe.pro"):
        entrar(c, u, "clave-segura-123")
        assert c.get("/api/bitacora").status_code == 403
        assert c.get("/api/bitacora/verificar").status_code == 403
        assert c.get("/api/reportes/bitacora.xlsx").status_code == 403
    assert c.get("/api/reportes/general.xlsx").status_code == 200  # el jefe descarga solo su área


def test_reabrir_carga_aprobada(client):
    c = client
    entrar(c, "gerente", "clave-segura-123")
    aprobada = next(m for m in c.get("/api/mediciones?estado=aprobado").json())
    assert c.post(f"/api/mediciones/{aprobada['id']}/reabrir", json={"accion": "observar"}).status_code == 422
    r = c.post(f"/api/mediciones/{aprobada['id']}/reabrir", json={"accion": "observar", "nota": "Dato digitado mal"})
    assert r.status_code == 200 and r.json()["estado"] == "observado"
    entrar(c, "jefe.sis", "otra-clave-123")
    assert c.post(f"/api/mediciones/{aprobada['id']}/reabrir", json={"accion": "observar", "nota": "no puedo"}).status_code == 403


def test_migracion_agrega_columnas_sin_perder_datos(tmp_path):
    import sqlite3, subprocess, sys, os
    db = tmp_path / "viejo.db"
    cx = sqlite3.connect(db)
    cx.executescript("""CREATE TABLE areas (id INTEGER PRIMARY KEY, nombre VARCHAR(80) UNIQUE);
        INSERT INTO areas (nombre) VALUES ('Existente');
        CREATE TABLE indicadores (id INTEGER PRIMARY KEY, codigo VARCHAR(20), nombre VARCHAR(150), objetivo TEXT,
          area_id INTEGER, unidad VARCHAR(20), frecuencia VARCHAR(20), sentido VARCHAR(10), formula VARCHAR(300),
          variables JSON, meta FLOAT, limite FLOAT, activo BOOLEAN, creado_en DATETIME);
        INSERT INTO indicadores VALUES (1,'OLD-01','Viejo','',1,'%','mensual','mayor','a / b * 100',
          '[{"clave":"a","etiqueta":"A"},{"clave":"b","etiqueta":"B"}]',90,80,1,'2026-01-01 00:00:00');""")
    cx.commit(); cx.close()
    code = ("from app.bootstrap import preparar_base; from app.database import SessionLocal; "
            "from app.models import Indicador; preparar_base(); "
            "db = SessionLocal(); i = db.get(Indicador, 1); print(i.codigo, i.decimales, i.constantes)")
    env = {**os.environ, "SIGT_DATABASE_URL": f"sqlite:///{db}", "SIGT_DATA_DIR": str(tmp_path)}
    out = subprocess.run([sys.executable, "-c", code], env=env, capture_output=True, text=True, cwd=os.getcwd())
    assert out.returncode == 0, out.stderr
    assert out.stdout.strip() == "OLD-01 None None"
