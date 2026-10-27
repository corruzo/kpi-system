"""Orden de los Excel exportados: nada superpuesto, nada recortado y filas/columnas del tamaño de su contenido."""
import math
from io import BytesIO

import pytest
from conftest import entrar, periodo_actual
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.excel.comun import altura_necesaria


@pytest.fixture(scope="module")
def libros(client):
    from conftest import crear_area, crear_indicador, crear_usuario
    area = crear_area(client)
    jefe = crear_usuario(client, "jefe", area)
    ind = crear_indicador(client, area, nombre="Indicador con un nombre bastante largo para obligar a que el texto se ajuste en varias líneas",
                          formula="(dias * 100) / tickets", variables=[{"clave": "dias", "etiqueta": "Días acumulados de atención al usuario"},
                                                                       {"clave": "tickets", "etiqueta": "Tickets"}])
    entrar(client, jefe)
    client.post("/api/mediciones", json={"indicador_id": ind["id"], "periodo": periodo_actual(client, ind["id"]),
                                         "valores": {"dias": 10, "tickets": 5}, "analisis": "A" * 300, "plan_mejora": "B " * 150})
    entrar(client, "admin", "clave-admin-123")
    rutas = {"general": "/api/reportes/general.xlsx", "fichas": f"/api/reportes/fichas.xlsx?indicador_id={ind['id']}",
             "bitacora": "/api/reportes/bitacora.xlsx"}
    return {k: load_workbook(BytesIO(client.get(r).content)) for k, r in rutas.items()}


def _hojas(libros):
    return [(f"{k}/{ws.title}", ws) for k, wb in libros.items() for ws in wb.worksheets]


def _alto_filas(ws, r1, r2):
    return sum(ws.row_dimensions[r].height or 15 for r in range(r1, r2 + 1))


def test_ningun_grafico_cubre_celdas_con_datos(libros):
    for nombre, ws in _hojas(libros):
        for g in ws._charts:
            ancla = g.anchor._from
            fila0, col0 = ancla.row + 1, ancla.col + 1
            alto_pt, ancho_cm = g.height / 0.0352778, g.width
            fila, acumulado = fila0, 0.0
            while acumulado < alto_pt:
                acumulado += ws.row_dimensions[fila].height or 15
                fila += 1
            col, acumulado = col0, 0.0
            while acumulado < ancho_cm:
                acumulado += ((ws.column_dimensions[get_column_letter(col)].width or 8.43) * 7 + 5) / 37.8
                col += 1
            cubiertas = [c.coordinate for r in ws.iter_rows(min_row=fila0, max_row=fila - 1, min_col=col0, max_col=col - 1)
                         for c in r if c.value not in (None, "")]
            assert not cubiertas, f"{nombre}: el gráfico cubre {cubiertas[:5]}"


def test_las_filas_con_texto_ajustado_tienen_la_altura_necesaria(libros):
    for nombre, ws in _hojas(libros):
        combinadas = {(rg.min_row, rg.min_col): rg for rg in ws.merged_cells.ranges}
        for fila in ws.iter_rows():
            for c in fila:
                if c.value in (None, "") or not (c.alignment and c.alignment.wrap_text):
                    continue
                rg = combinadas.get((c.row, c.column))
                ancho = sum(ws.column_dimensions[get_column_letter(k)].width or 8.43
                            for k in range(rg.min_col, rg.max_col + 1)) if rg else ws.column_dimensions[c.column_letter].width or 8.43
                alto = _alto_filas(ws, c.row, rg.max_row if rg else c.row)
                assert alto >= altura_necesaria(c, ancho) - 1, f"{nombre}!{c.coordinate} ({alto} < {altura_necesaria(c, ancho):.0f})"


def test_el_texto_sin_ajuste_cabe_en_su_columna(libros):
    for nombre, ws in _hojas(libros):
        combinadas = {(r, c) for rg in ws.merged_cells.ranges for r in range(rg.min_row, rg.max_row + 1)
                      for c in range(rg.min_col, rg.max_col + 1)}
        for fila in ws.iter_rows():
            if sum(1 for c in fila if c.value not in (None, "")) < 2:
                continue                                      # títulos de sección: se extienden sobre celdas vacías
            for c in fila:
                if not isinstance(c.value, str) or c.value.startswith("=") or (c.row, c.column) in combinadas:
                    continue
                if c.alignment and c.alignment.wrap_text:
                    continue
                ancho = ws.column_dimensions[c.column_letter].width or 8.43
                assert max(len(t) for t in c.value.split("\n")) <= math.ceil(ancho * 1.25), f"{nombre}!{c.coordinate}: «{c.value[:40]}»"


def test_la_ficha_reserva_un_espacio_propio_para_el_logo_o_el_nombre(libros):
    ws = libros["fichas"].worksheets[0]
    assert ws["D1"].value == "FICHA TÉCNICA DE INDICADOR DE GESTIÓN"
    assert any(str(rg) in ("A1:C2", "B1:C2") for rg in ws.merged_cells.ranges)
