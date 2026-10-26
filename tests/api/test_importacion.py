import pathlib
from io import BytesIO

from conftest import entrar
from openpyxl import load_workbook

from app.dominio import calendario

LIBRO = pathlib.Path(__file__).parent.parent / "datos" / "fichas_historicas.xlsx"


def test_importar_fichas_historicas_y_reemplazar_el_excel(client):
    """Libro de 5 hojas con ejercicio octubre–septiembre: previsualiza, confirma y genera la ficha del SGC."""
    entrar(client, "admin", "clave-admin-123")
    client.put("/api/parametros", json={"empresa": "TRIPOLIVEN, C.A.", "gerencia_emisora": "Gerencia de Normalización y Calidad",
                                        "codigo_formato": "NC-16-7", "ejercicio_inicio": 10})
    r = client.post("/api/importar/excel", files={"archivo": ("fichas.xlsx", LIBRO.read_bytes())})
    assert r.status_code == 200, r.text
    prev = r.json()
    assert len(prev["indicadores"]) == 5 and prev["inicio_ejercicio"] == 10
    sp = next(i for i in prev["indicadores"] if i["sigla"] == "SP")
    assert sp["formula"] == "sg*100/sr" and sp["frecuencia"] == "mensual" and len(sp["mediciones"]) == 8
    assert sp["mediciones"][0]["periodo"] == "2025-10" and sp["mediciones"][0]["resultado"] == 92.86
    io = next(i for i in prev["indicadores"] if i["sigla"] == "IO")
    assert io["frecuencia"] == "trimestral" and [m["periodo"] for m in io["mediciones"]] == ["2026-Q1", "2026-Q2"]

    cuerpo = {"archivo": "fichas.xlsx", "inicio_ejercicio": prev["inicio_ejercicio"], "indicadores": prev["indicadores"],
              "areas_nuevas": ["Sistemas"], "catalogos": [{"tipo": "gerencia", "valor": g} for g in prev["catalogos"]["gerencia"]]}
    ok = client.post("/api/importar/confirmar", json=cuerpo)
    if calendario.inicio_mes() != 10:                      # el calendario del sistema debe coincidir con el de las fichas
        assert ok.status_code == 409
        return
    assert ok.status_code == 201, ok.text
    assert ok.json()["indicadores"] == 5 and ok.json()["mediciones"] == 8 + 8 + 8 + 2 + 7
    assert client.post("/api/importar/confirmar", json=cuerpo).status_code == 422       # códigos ya existentes

    ind = next(i for i in client.get("/api/indicadores").json() if i["sigla"] == "SP")
    f = client.get(f"/api/indicadores/{ind['id']}/ficha?ejercicio=2026").json()
    assert f["ejercicio_etiqueta"] == "2025-2026" and len(f["periodos"]) == 12 and f["periodos"][0]["corta"] == "Oct."
    assert f["periodos"][0]["medicion"]["resultado"] == 92.86 and f["indicador"]["meta_texto"] == "96 Mín."
    x = client.get("/api/reportes/fichas.xlsx?ejercicio=2026")
    assert x.status_code == 200
    libro = load_workbook(BytesIO(x.content))
    assert len(libro.sheetnames) >= 5 and libro[libro.sheetnames[0]]["D1"].value == "FICHA TÉCNICA DE INDICADOR DE GESTIÓN"
