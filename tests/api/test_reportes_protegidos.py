from io import BytesIO

import pytest
from conftest import entrar, periodo_actual
from openpyxl import load_workbook

from app.dominio import calendario


def configurar_proteccion(client, activa: bool, clave: str = ""):
    entrar(client, "admin", "clave-admin-123")
    r = client.put("/api/parametros", json={"empresa": "TRIPOLIVEN, C.A.", "gerencia_emisora": "", "codigo_formato": "",
                                            "ejercicio_inicio": calendario.inicio_mes(), "proteger_reportes": activa,
                                            "clave_reportes": clave})
    return r


@pytest.fixture()
def reporte_protegido(client, escenario):
    """Un indicador con una carga, y la protección de reportes activada (se desactiva al terminar)."""
    ind = escenario["indicador"]
    entrar(client, escenario["jefe"])
    r = client.post("/api/mediciones", json={"indicador_id": ind["id"], "periodo": periodo_actual(client, ind["id"]),
                                             "valores": {"dias": 10, "tickets": 5}})
    assert r.status_code == 201
    assert configurar_proteccion(client, True, "secreta-123").status_code == 200
    yield ind
    configurar_proteccion(client, False)


def descargar_fichas(client, ind):
    r = client.get(f"/api/reportes/fichas.xlsx?indicador_id={ind['id']}")
    assert r.status_code == 200
    return r.content


def test_la_clave_no_se_expone_y_se_exige_al_activar(client):
    assert configurar_proteccion(client, True).status_code == 422                 # sin contraseña
    assert configurar_proteccion(client, True, "corta").status_code == 422
    r = configurar_proteccion(client, True, "secreta-123")
    try:
        assert r.status_code == 200
        assert r.json()["proteger_reportes"] is True and r.json()["clave_reportes_definida"] is True
        assert "secreta-123" not in str(client.get("/api/parametros").json())
        assert configurar_proteccion(client, True).status_code == 200              # en blanco = conserva la actual
    finally:
        configurar_proteccion(client, False)


def test_la_ficha_exportada_tiene_formulas_y_hojas_bloqueadas(client, reporte_protegido):
    libro = load_workbook(BytesIO(descargar_fichas(client, reporte_protegido)))
    assert libro.security.lockStructure
    assert all(ws.protection.sheet for ws in libro.worksheets)
    hoja = libro.worksheets[0]
    formulas = [c.value for fila in hoja.iter_rows() for c in fila if isinstance(c.value, str) and c.value.startswith("=")]
    assert any(f.startswith("=ROUND(") and "/" in f for f in formulas)             # resultado: dias / tickets
    assert any(f.startswith("=IF(") and "Cumple" in f for f in formulas)           # semáforo
    assert "Control" in libro.sheetnames


def test_sin_proteccion_las_hojas_quedan_abiertas_pero_sellado(client, reporte_protegido):
    configurar_proteccion(client, False)
    libro = load_workbook(BytesIO(descargar_fichas(client, reporte_protegido)))
    assert not any(ws.protection.sheet for ws in libro.worksheets) and "Control" in libro.sheetnames


def test_todos_los_reportes_se_bloquean(client, reporte_protegido):
    for ruta in ("/api/reportes/general.xlsx", "/api/reportes/bitacora.xlsx"):
        libro = load_workbook(BytesIO(client.get(ruta).content))
        assert all(ws.protection.sheet for ws in libro.worksheets), ruta


def test_verificacion_detecta_cambios_en_el_contenido(client, reporte_protegido):
    original = descargar_fichas(client, reporte_protegido)
    ok = client.post("/api/reportes/verificar", files={"archivo": ("r.xlsx", original)})
    assert ok.status_code == 200 and ok.json()["ok"] is True

    libro = load_workbook(BytesIO(original))
    celda = next(c for fila in libro.worksheets[0].iter_rows() for c in fila if isinstance(c.value, (int, float)))
    celda.value += 1                                                                # alguien altera un dato
    salida = BytesIO()
    libro.save(salida)
    alterado = client.post("/api/reportes/verificar", files={"archivo": ("r.xlsx", salida.getvalue())}).json()
    assert alterado["ok"] is False and "modificado" in alterado["motivo"]

    libro = load_workbook(BytesIO(original))
    libro["Control"]["B4"] = "R-20200101-ABCDEF"                                    # número inventado
    salida = BytesIO()
    libro.save(salida)
    assert client.post("/api/reportes/verificar", files={"archivo": ("r.xlsx", salida.getvalue())}).json()["ok"] is False


def test_verificar_rechaza_archivos_ajenos_y_a_los_jefes(client, escenario):
    from openpyxl import Workbook
    vacio = BytesIO()
    Workbook().save(vacio)
    entrar(client, escenario["gerente"])
    assert client.post("/api/reportes/verificar", files={"archivo": ("x.xlsx", vacio.getvalue())}).status_code == 422
    entrar(client, escenario["jefe"])
    assert client.post("/api/reportes/verificar", files={"archivo": ("x.xlsx", vacio.getvalue())}).status_code == 403
