from io import BytesIO

import pytest
from conftest import crear_area, crear_usuario, entrar
from openpyxl import load_workbook


def test_la_bitacora_es_inalterable_y_verificable(client):
    from app.database import SessionLocal
    from app.models import Bitacora
    entrar(client, "admin", "clave-admin-123")
    assert client.get("/api/bitacora/verificar").json()["ok"] is True
    with SessionLocal() as db:
        registro = db.query(Bitacora).first()
        registro.accion = "X"
        with pytest.raises(RuntimeError):
            db.flush()


@pytest.mark.parametrize("rol", ["jefe", "gerente"])
def test_solo_el_administrador_ve_la_bitacora(client, rol):
    entrar(client, crear_usuario(client, rol, crear_area(client) if rol == "jefe" else None))
    for ruta in ("/api/bitacora", "/api/bitacora/verificar", "/api/reportes/bitacora.xlsx"):
        assert client.get(ruta).status_code == 403


def test_reporte_general_en_excel(client, escenario):
    entrar(client, "admin", "clave-admin-123")
    r = client.get("/api/reportes/general.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    libro = load_workbook(BytesIO(r.content))
    assert libro.sheetnames == ["Resumen", "Mediciones", "Tendencias", "Plan de mejora", "Fichas técnicas"]
    assert client.get("/api/reportes/bitacora.xlsx").status_code == 200
    entrar(client, escenario["jefe"])
    assert client.get("/api/reportes/general.xlsx").status_code == 200       # el jefe descarga solo su área
