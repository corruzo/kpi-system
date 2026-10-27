"""Reporte general en Excel: resumen con gráfico, mediciones, tendencias, plan de mejora y fichas."""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..dominio import calendario
from ..dominio.ficha_texto import formula_en_palabras
from .comun import (
    BORDE,
    BRAND,
    ESTADO,
    FUENTE,
    GRIS,
    NAVY,
    SEM,
    ZEBRA,
    ajustar,
    celda_semaforo,
    configurar_impresion,
    encabezado_tabla,
    formato_numero,
    franja_titulo,
    hora_local,
    relleno,
    ubicar_grafico,
)
from .sello import colorear_semaforo, formula_semaforo


def reporte_general(
    mediciones, indicadores, generado_por: str, filtros: str, por_aprobar: int, logo: bytes | None = None
) -> Workbook:
    """mediciones: [Medicion]; indicadores: [Indicador] (los incluidos en el reporte)."""
    wb = Workbook()
    ahora = hora_local(datetime.utcnow())
    sub = f"Reporte de indicadores de gestión · Generado el {ahora:%d/%m/%Y %H:%M} por {generado_por}"

    # ---------- Resumen
    ws = wb.active
    ws.title = "Resumen"
    franja_titulo(ws, "SGC • TRIPOLIVEN — Indicadores de gestión", sub, 7, logo)
    ws["A4"], ws["B4"] = "Filtros aplicados", filtros or "Ninguno (todos los datos)"
    ws["A4"].font = Font(name=FUENTE, bold=True, color=GRIS)
    ws.merge_cells("B4:G4")
    ws["B4"].alignment = Alignment(wrap_text=True, vertical="center")
    ultimo = {}
    for m in sorted(mediciones, key=lambda m: (m.indicador_id, m.periodo)):
        if m.estado in ("enviado", "aprobado"):
            ultimo[m.indicador_id] = m
    conteo = {"verde": 0, "amarillo": 0, "rojo": 0}
    por_area: dict[str, dict] = {}
    for i in indicadores:
        a = por_area.setdefault(i.area.nombre, {"total": 0, "verde": 0, "amarillo": 0, "rojo": 0, "sin": 0})
        a["total"] += 1
        m = ultimo.get(i.id)
        if m:
            conteo[m.semaforo] += 1
            a[m.semaforo] += 1
        else:
            a["sin"] += 1
    total = len(indicadores)
    kpis = [
        ("Indicadores", total),
        ("Cumplimiento (verde)", conteo["verde"] / total if total else 0),
        ("En alerta", conteo["amarillo"]),
        ("Críticos", conteo["rojo"]),
        ("Cargas registradas", len(mediciones)),
        ("Por aprobar", por_aprobar),
    ]
    colores = [NAVY, "16A34A", "F59E0B", "DC2626", BRAND, GRIS]
    for k, ((t, v), col) in enumerate(zip(kpis, colores, strict=False)):
        r, c = 6 + (k // 3) * 3, 1 + (k % 3) * 2
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 1)
        ws.merge_cells(start_row=r + 1, start_column=c, end_row=r + 1, end_column=c + 1)
        a = ws.cell(r, c, t.upper())
        a.font = Font(name=FUENTE, size=9, bold=True, color="FFFFFF")
        a.fill = relleno(col)
        a.alignment = Alignment(horizontal="center")
        ws.cell(r, c + 1).fill = relleno(col)
        b = ws.cell(r + 1, c, v)
        b.font = Font(name=FUENTE, size=20, bold=True, color=col)
        b.alignment = Alignment(horizontal="center", vertical="center")
        b.fill = relleno("FFFFFF")
        b.border = BORDE
        ws.cell(r + 1, c + 1).border = BORDE
        if t.startswith("Cumpl"):
            b.number_format = "0.0%"
        ws.row_dimensions[r + 1].height = 34
    f0 = 13
    ws.cell(f0 - 1, 1, "Cumplimiento por departamento (última medición de cada indicador)").font = Font(
        name=FUENTE, bold=True, size=12, color=NAVY
    )
    encabezado_tabla(
        ws,
        f0,
        ["Área", "Indicadores", "Cumple", "Alerta", "Crítico", "Sin datos", "% cumplimiento"],
        [30, 16, 14, 14, 14, 14, 20],
    )
    for k, (nombre, a) in enumerate(sorted(por_area.items()), f0 + 1):
        fila = [
            nombre,
            a["total"],
            a["verde"],
            a["amarillo"],
            a["rojo"],
            a["sin"],
            f"=IF(B{k}=0,0,C{k}/B{k})",
        ]
        for j, v in enumerate(fila, 1):
            c = ws.cell(k, j, v)
            c.font = Font(name=FUENTE, size=10, bold=(j == 1))
            c.border = BORDE
            c.alignment = Alignment(horizontal="left" if j == 1 else "center")
            if k % 2:
                c.fill = relleno(ZEBRA)
        ws.cell(k, 7).number_format = "0.0%"
    fin = f0 + len(por_area)
    if por_area:
        fila_total = fin + 1
        ws.cell(fila_total, 1, "Total")
        for j in range(2, 7):
            ws.cell(fila_total, j, f"=SUM({get_column_letter(j)}{f0 + 1}:{get_column_letter(j)}{fin})")
        ws.cell(fila_total, 7, f"=IF(B{fila_total}=0,0,C{fila_total}/B{fila_total})").number_format = "0.0%"
        for j in range(1, 8):
            c = ws.cell(fila_total, j)
            c.font, c.border = Font(name=FUENTE, size=10, bold=True), BORDE
            c.alignment = Alignment(horizontal="left" if j == 1 else "center")
            c.fill = relleno("E2E8F0")
        ws.conditional_formatting.add(
            f"G{f0 + 1}:G{fin}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="16A34A")
        )
        ch = BarChart()
        ch.type, ch.grouping, ch.overlap = "col", "stacked", 100
        ch.title = "Semáforos por departamento"
        ch.add_data(Reference(ws, min_col=3, max_col=5, min_row=f0, max_row=fin), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=f0 + 1, max_row=fin))
        for s, col in zip(ch.series, ("16A34A", "F59E0B", "DC2626"), strict=False):
            s.graphicalProperties.solidFill = col
            s.graphicalProperties.line.solidFill = col
        ch.x_axis.delete = ch.y_axis.delete = False  # Excel oculta los ejes si no se indica
        ch.y_axis.numFmt, ch.y_axis.majorGridlines = "0", None
        ubicar_grafico(ch, ws, fin + 4, 20, 8.5)
    configurar_impresion(ws, f0, apaisado=False)
    ajustar(ws)

    # ---------- Mediciones
    ws = wb.create_sheet("Mediciones")
    cols = [
        "Código",
        "Indicador",
        "Área",
        "Período",
        "Ejercicio fiscal",
        "Frecuencia",
        "Resultado",
        "Meta",
        "Límite crítico",
        "Sentido",
        "Semáforo",
        "Estado",
        "Cargado por",
        "Fecha de carga",
        "Revisado por",
        "Fecha de revisión",
        "Nota de revisión",
        "Análisis de resultados",
        "Correcciones / plan de mejora",
        "Versión",
    ]
    anchos = [10, 38, 20, 16, 12, 12, 14, 12, 14, 12, 12, 13, 24, 17, 24, 17, 34, 44, 44, 9]
    franja_titulo(ws, "Mediciones registradas", sub, len(cols))
    encabezado_tabla(ws, 4, cols, anchos)
    ordenadas = sorted(mediciones, key=lambda m: (m.indicador.codigo, m.periodo))
    for k, m in enumerate(ordenadas, 5):
        i = m.indicador
        dec = m.decimales if m.decimales is not None else (i.decimales if i.decimales is not None else 2)
        fila = [
            i.codigo,
            i.nombre,
            i.area.nombre,
            calendario.etiqueta(i.frecuencia, m.periodo),
            calendario.etiqueta_ejercicio(calendario.ejercicio_de(i.frecuencia, m.periodo)),
            i.frecuencia.capitalize(),
            m.resultado,
            m.meta,
            m.limite,
            "Mayor es mejor" if m.sentido == "mayor" else "Menor es mejor",
            None,
            ESTADO[m.estado],
            m.cargado_por.nombre,
            hora_local(m.cargado_en),
            m.revisado_por.nombre if m.revisado_por else "",
            hora_local(m.revisado_en),
            m.nota_revision,
            m.analisis,
            m.plan_mejora,
            m.version,
        ]
        for j, v in enumerate(fila, 1):
            c = ws.cell(k, j, v)
            c.font = Font(name=FUENTE, size=10)
            c.border = BORDE
            c.alignment = Alignment(vertical="center", wrap_text=j in (2, 17, 18, 19))
            if k % 2:
                c.fill = relleno(ZEBRA)
        for j in (7, 8, 9):
            ws.cell(k, j).number_format = formato_numero(i.unidad, dec)
        ws.cell(k, 7).font = Font(name=FUENTE, size=10, bold=True)
        for j in (14, 16):
            ws.cell(k, j).number_format = "dd/mm/yyyy hh:mm"
        if m.meta_tope is None and m.limite_tope is None:
            ws.cell(k, 11).value = formula_semaforo(f"G{k}", f"H{k}", f"I{k}", m.sentido == "mayor")
            ws.cell(k, 11).alignment = Alignment(horizontal="center", vertical="center")
        else:                                   # con topes, la regla no cabe en una fórmula simple
            celda_semaforo(ws.cell(k, 11), m.semaforo)
        ws.cell(k, 11).border = BORDE
    colorear_semaforo(ws, f"K5:K{max(5, 4 + len(ordenadas))}")
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(5, 4 + len(ordenadas))}"
    configurar_impresion(ws, 4)
    ajustar(ws)

    # ---------- Tendencias (matriz indicador × período, por frecuencia)
    ws = wb.create_sheet("Tendencias")
    franja_titulo(ws, "Tendencias por indicador", sub, 12)
    fila = 4
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 38
    for freq in ("mensual", "bimestral", "trimestral", "cuatrimestral", "semestral", "anual", "eventual"):
        inds = [i for i in indicadores if i.frecuencia == freq]
        if not inds:
            continue
        validas = [m for m in mediciones if m.indicador.frecuencia == freq and m.estado in ("enviado", "aprobado")]
        periodos = sorted({m.periodo for m in validas})
        ws.cell(fila, 1, f"Indicadores {freq}es" if freq != "anual" else "Indicadores anuales").font = Font(
            name=FUENTE, bold=True, size=12, color=NAVY
        )
        fila += 1
        encabezado_tabla(ws, fila, ["Código", "Indicador"] + [calendario.etiqueta(freq, p) for p in periodos] + ["Meta"])
        for j in range(3, 3 + len(periodos) + 1):
            ws.column_dimensions[get_column_letter(j)].width = max(ws.column_dimensions[get_column_letter(j)].width or 0, 15)
        por = {(m.indicador_id, m.periodo): m for m in validas}
        for i in sorted(inds, key=lambda x: x.codigo):
            fila += 1
            dec = i.decimales if i.decimales is not None else 2
            ws.cell(fila, 1, i.codigo).font = Font(name=FUENTE, bold=True, size=10)
            ws.cell(fila, 1).alignment = Alignment(vertical="center")
            ws.cell(fila, 2, i.nombre).alignment = Alignment(wrap_text=True, vertical="center")
            ws.cell(fila, 2).font = Font(name=FUENTE, size=10)
            for j, p in enumerate(periodos, 3):
                m = por.get((i.id, p))
                c = ws.cell(fila, j, m.resultado if m else None)
                c.number_format = formato_numero(i.unidad, dec)
                c.alignment = Alignment(horizontal="center", vertical="center")
                if m:
                    bg, fg, _ = SEM[m.semaforo]
                    c.fill = relleno(bg)
                    c.font = Font(name=FUENTE, bold=True, color=fg, size=10)
            c = ws.cell(fila, 3 + len(periodos), i.meta)
            c.number_format = formato_numero(i.unidad, dec)
            c.font = Font(name=FUENTE, italic=True, color=GRIS, size=10)
            c.alignment = Alignment(horizontal="center", vertical="center")
            for j in range(1, 4 + len(periodos)):
                ws.cell(fila, j).border = BORDE
        fila += 3
    ws.freeze_panes = "C4"
    configurar_impresion(ws, 4)
    ajustar(ws)

    # ---------- Plan de mejora (ISO 9001: 9.1.3 y 10.2)
    ws = wb.create_sheet("Plan de mejora")
    cols = [
        "Código",
        "Indicador",
        "Área",
        "Período",
        "Ejercicio fiscal",
        "Resultado",
        "Meta",
        "Semáforo",
        "Estado",
        "Análisis de resultados",
        "Correcciones / plan de mejora",
        "Responsable de la carga",
    ]
    franja_titulo(ws, "Desviaciones y acciones correctivas (ISO 9001:2015, 9.1.3 y 10.2)", sub, len(cols))
    encabezado_tabla(ws, 4, cols, [10, 36, 20, 16, 12, 13, 12, 12, 13, 50, 50, 24])
    desv = [m for m in ordenadas if m.semaforo != "verde"]
    for k, m in enumerate(sorted(desv, key=lambda m: (m.semaforo != "rojo", m.indicador.codigo, m.periodo)), 5):
        i = m.indicador
        dec = m.decimales if m.decimales is not None else 2
        fila = [
            i.codigo,
            i.nombre,
            i.area.nombre,
            calendario.etiqueta(i.frecuencia, m.periodo),
            calendario.etiqueta_ejercicio(calendario.ejercicio_de(i.frecuencia, m.periodo)),
            m.resultado,
            m.meta,
            None,
            ESTADO[m.estado],
            m.analisis,
            m.plan_mejora,
            m.cargado_por.nombre,
        ]
        for j, v in enumerate(fila, 1):
            c = ws.cell(k, j, v)
            c.font = Font(name=FUENTE, size=10)
            c.border = BORDE
            c.alignment = Alignment(vertical="center", wrap_text=j in (2, 10, 11))
        for j in (6, 7):
            ws.cell(k, j).number_format = formato_numero(i.unidad, dec)
        celda_semaforo(ws.cell(k, 8), m.semaforo)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(5, 4 + len(desv))}"
    configurar_impresion(ws, 4)
    ajustar(ws)

    # ---------- Fichas técnicas
    ws = wb.create_sheet("Fichas técnicas")
    cols = [
        "Código",
        "Indicador",
        "Área",
        "Seguimiento / proceso",
        "Líder de proceso",
        "Resp. medición",
        "Objetivo",
        "Fuente de datos",
        "Unidad",
        "Frecuencia",
        "Sentido",
        "Meta (verde)",
        "Límite crítico",
        "Variables",
        "Constantes",
        "Fórmula",
        "Fórmula (en palabras)",
        "Método / notas",
        "Estado",
    ]
    franja_titulo(ws, "Fichas técnicas de los indicadores", sub, len(cols))
    encabezado_tabla(ws, 4, cols, [10, 36, 20, 34, 22, 22, 40, 28, 10, 13, 15, 12, 13, 40, 28, 34, 46, 40, 11])
    for k, i in enumerate(sorted(indicadores, key=lambda x: x.codigo), 5):
        consts = i.constantes or []
        fila = [
            i.codigo,
            i.nombre,
            i.area.nombre,
            i.ref_nombre or "",
            i.lider or "",
            i.resp_medicion or "",
            i.objetivo,
            i.fuente or "",
            i.unidad,
            i.frecuencia.capitalize(),
            "Mayor es mejor" if i.sentido == "mayor" else "Menor es mejor",
            i.meta,
            i.limite,
            "\n".join(f"• {v['etiqueta']} ({v['clave']})" for v in i.variables),
            "\n".join(f"• {c['etiqueta']} = {c['valor']}" for c in consts),
            i.formula,
            formula_en_palabras(i.formula, i.variables, consts),
            i.notas or "",
            "Activo" if i.activo else "Archivado",
        ]
        for j, v in enumerate(fila, 1):
            c = ws.cell(k, j, v)
            c.font = Font(name="Consolas" if j == 16 else FUENTE, size=10)
            c.border = BORDE
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if k % 2:
                c.fill = relleno(ZEBRA)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(5, 4 + len(indicadores))}"
    configurar_impresion(ws, 4)
    ajustar(ws)
    return wb


# ------------------------------------------------------------------ bitácora
