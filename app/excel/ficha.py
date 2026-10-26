"""Ficha técnica de indicador en Excel, con el formato del SGC (una hoja por indicador)."""

import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..dominio import calendario
from ..dominio.ficha_texto import formula_documentada, franjas, meta_texto
from ..dominio.formula_excel import a_excel
from .comun import (
    BORDE,
    BRAND,
    FUENTE,
    GRIS,
    NAVY,
    SEM,
    celda_semaforo,
    formato_numero,
    hora_local,
    poner_logo,
    relleno,
)
from .sello import colorear_semaforo, formula_semaforo


def _banda(ws, fila, texto, ncols=13):
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ncols)
    c = ws.cell(fila, 1, texto)
    c.font = Font(name=FUENTE, bold=True, size=10, color="FFFFFF")
    c.fill = relleno(NAVY)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[fila].height = 20


def _par(ws, fila, c1, etiqueta, c2, c3, valor, alto=None, mono=False, wrap=True):
    """Celda de etiqueta (c1) + valor combinado (c2..c3)."""
    e = ws.cell(fila, c1, etiqueta)
    e.font = Font(name=FUENTE, bold=True, size=10, color=NAVY)
    e.fill = relleno("E2E8F0")
    e.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    e.border = BORDE
    if c3 > c2:
        ws.merge_cells(start_row=fila, start_column=c2, end_row=fila, end_column=c3)
    v = ws.cell(fila, c2, valor)
    v.font = Font(name="Consolas" if mono else FUENTE, size=10)
    v.alignment = Alignment(vertical="top" if mono else "center", wrap_text=wrap, indent=1)
    for c in range(c2, c3 + 1):
        ws.cell(fila, c).border = BORDE
    if alto:
        ws.row_dimensions[fila].height = alto
    return v


def _lineas(texto: str, ancho: int) -> int:
    return sum(max(1, -(-len(t) // ancho)) for t in (texto or "").split("\n"))


def fichas_xlsx(fichas, params: dict, generado_por: str) -> Workbook:
    """fichas: [(Indicador, ejercicio, periodos[{periodo, etiqueta, corta, medicion}])]. Una hoja por indicador."""
    wb = Workbook()
    wb.remove(wb.active)
    usados: set[str] = set()
    ahora = hora_local(datetime.utcnow())
    for ind, y, periodos in fichas:
        base = re.sub(r"[\\/*?:\[\]]", "", ind.sigla or ind.codigo)[:28] or ind.codigo
        nombre, n = base, 2
        while nombre.lower() in usados:
            nombre, n = f"{base[:26]}_{n}", n + 1
        usados.add(nombre.lower())
        ws = wb.create_sheet(nombre)
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 30
        for j in range(2, 14):
            ws.column_dimensions[get_column_letter(j)].width = 11.5
        tipo = {
            "proceso": "Desempeño de proceso",
            "objetivo": "Objetivo organizacional",
            "proveedor": "Desempeño de proveedor externo",
        }.get(ind.seguimiento or "proceso", "")
        etq1, etq2 = {
            "objetivo": ("OBJETIVO ORGANIZACIONAL", "OBJETIVO ESPECÍFICO"),
            "proveedor": ("PROCESO", "BIEN o SERVICIO"),
        }.get(ind.seguimiento or "proceso", ("PROCESO", "OBJETIVO DEL PROCESO"))
        dec = ind.decimales if ind.decimales is not None else 2
        fmt = formato_numero(ind.unidad, dec)

        # --- encabezado
        ws.merge_cells("A1:C2")
        c = ws["A1"]
        c.value = params.get("empresa", "")
        c.font = Font(name=FUENTE, bold=True, size=14, color="FFFFFF")
        c.fill = relleno(NAVY)
        c.alignment = Alignment(vertical="center", indent=1, wrap_text=True)
        ws.merge_cells("D1:J2")
        c = ws["D1"]
        c.value = "FICHA TÉCNICA DE INDICADOR DE GESTIÓN"
        c.font = Font(name=FUENTE, bold=True, size=14, color="FFFFFF")
        c.fill = relleno(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells("K1:M1")
        ws["K1"].value = "Seguimiento de:"
        ws.merge_cells("K2:M2")
        ws["K2"].value = tipo
        for ref, bold in (("K1", False), ("K2", True)):
            ws[ref].font = Font(name=FUENTE, size=10, bold=bold, color="FFFFFF")
            ws[ref].fill = relleno(BRAND)
            ws[ref].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = ws.row_dimensions[2].height = 20
        if poner_logo(ws, params.get("logo_bytes"), "A1", 38):
            c.alignment = Alignment(vertical="center", indent=1)
            ws["A1"].alignment = Alignment(vertical="center", indent=7, wrap_text=True)
        ws.merge_cells("A3:M3")
        ws["A3"].value = params.get("gerencia_emisora", "")
        ws["A3"].font = Font(name=FUENTE, italic=True, size=10, color=GRIS)
        ws["A3"].alignment = Alignment(indent=1)

        # --- proceso / objetivo
        _banda(ws, 4, etq1)
        ws.merge_cells("A5:M5")
        ws["A5"].value = ind.ref_nombre or ""
        ws["A5"].font = Font(name=FUENTE, bold=True, size=11)
        ws["A5"].alignment = Alignment(indent=1, vertical="center")
        _banda(ws, 6, etq2)
        ws.merge_cells("A7:M7")
        ws["A7"].value = ind.ref_detalle or ind.objetivo or ""
        ws["A7"].font = Font(name=FUENTE, size=10)
        ws["A7"].alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.row_dimensions[7].height = max(18, 14 * _lineas(ws["A7"].value, 150))

        # --- definición + semáforo
        _banda(ws, 8, "DEFINICIÓN DEL INDICADOR E INFORMACIÓN PARA LA MEDICIÓN")
        nombre_ind = ind.nombre + (f" ({ind.sigla})" if ind.sigla and f"({ind.sigla})" not in ind.nombre else "")
        _par(ws, 9, 1, "Nombre del indicador", 2, 7, nombre_ind, 30)
        _par(ws, 10, 1, "Unidad de medida", 2, 3, ind.unidad)
        _par(ws, 10, 4, "Meta", 5, 7, meta_texto(ind))
        _par(ws, 11, 1, "Frecuencia de medición", 2, 7, ind.frecuencia.capitalize())
        doc = formula_documentada(ind)
        _par(ws, 12, 1, "Fórmula de cálculo", 2, 7, doc, mono=True)
        ws.merge_cells(start_row=12, start_column=2, end_row=14, end_column=7)
        for r in (12, 13, 14):
            ws.row_dimensions[r].height = max(20, 14 * _lineas(doc, 60) / 3 + 4)
        ws.merge_cells("A12:A14")
        _par(ws, 15, 1, "Fuente de la información", 2, 7, ind.fuente or "", 22)
        # semáforo
        ws.merge_cells("I9:M9")
        c = ws["I9"]
        c.value = "SEMÁFORO"
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF", size=10)
        c.fill = relleno(BRAND)
        c.alignment = Alignment(horizontal="center", vertical="center")
        for col, t in ((9, ""), (10, "Mín."), (12, "Máx.")):
            if col == 9:
                continue
            ws.merge_cells(start_row=10, start_column=col, end_row=10, end_column=col + 1)
            h = ws.cell(10, col, t)
            h.font = Font(name=FUENTE, bold=True, size=9, color=GRIS)
            h.alignment = Alignment(horizontal="center")
        verde, amar, rojo = franjas(ind)
        for r, (nom, (mn, mx), clave) in enumerate(
            (("Verde", verde, "verde"), ("Amarillo", amar, "amarillo"), ("Rojo", rojo, "rojo")), 11
        ):
            bg, fg, _ = SEM[clave]
            e = ws.cell(r, 9, nom)
            e.fill, e.font = relleno(bg), Font(name=FUENTE, bold=True, size=10, color=fg)
            e.alignment = Alignment(horizontal="center", vertical="center")
            e.border = BORDE
            for col, val in ((10, mn), (12, mx)):
                ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)
                v = ws.cell(r, col, val if val is not None else "")
                v.number_format = fmt
                v.alignment = Alignment(horizontal="center", vertical="center")
                v.font = Font(name=FUENTE, size=10)
                for cc in (col, col + 1):
                    ws.cell(r, cc).border = BORDE
        _par(ws, 15, 9, "Ejercicio fiscal", 10, 13, calendario.etiqueta_ejercicio(y))

        # --- responsabilidades
        _banda(ws, 16, "RESPONSABILIDADES")
        _par(ws, 17, 1, "Gerencia", 2, 4, ind.gerencia or "", 22)
        _par(ws, 17, 5, "Departamento", 6, 8, ind.area.nombre)
        _par(ws, 17, 9, "Líder de proceso", 10, 13, ind.lider or "")
        _par(ws, 18, 1, "Responsable de medición", 2, 4, ind.resp_medicion or "", 22)
        _par(ws, 18, 5, "Responsable de registro", 6, 8, ind.resp_registro or "")
        _par(ws, 18, 9, "Responsable de análisis", 10, 13, ind.resp_analisis or "")

        # --- resultados de las determinaciones
        _banda(ws, 19, "RESULTADOS DE LAS DETERMINACIONES")
        hdr = ws.cell(20, 1, "PERÍODOS")
        hdr.font = Font(name=FUENTE, bold=True, size=10, color=NAVY)
        hdr.fill = relleno("E2E8F0")
        hdr.border = BORDE
        for j in range(12):
            c = ws.cell(20, 2 + j, periodos[j]["corta"] if j < len(periodos) else "No aplica")
            c.font = Font(name=FUENTE, bold=True, size=9, color="FFFFFF" if j < len(periodos) else GRIS)
            c.fill = relleno(BRAND if j < len(periodos) else "E2E8F0")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BORDE
        fila = 21
        fila_variable = {v["clave"]: fila + k for k, v in enumerate(ind.variables)}
        for v in ind.variables:
            e = ws.cell(fila, 1, v["etiqueta"])
            e.font = Font(name=FUENTE, size=10)
            e.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
            e.border = BORDE
            ws.row_dimensions[fila].height = max(18, 14 * _lineas(v["etiqueta"], 30))
            for j in range(12):
                c = ws.cell(fila, 2 + j)
                c.border = BORDE
                c.alignment = Alignment(horizontal="center", vertical="center")
                if j < len(periodos) and periodos[j]["medicion"] and periodos[j]["medicion"].estado != "observado":
                    c.value = periodos[j]["medicion"].valores.get(v["clave"])
                    if c.value is None and v.get("opcional"):
                        c.value = v.get("defecto", 0)          # el sistema usó el valor por defecto
                    c.number_format = "#,##0.##"
                elif j >= len(periodos):
                    c.fill = relleno("F1F5F9")
            fila += 1

        # --- valores del indicador
        fila += 1
        _banda(ws, fila, "VALORES DEL INDICADOR")
        fila += 1
        for j, t in enumerate(["Período", "Indicador", "Mín.", "Máx.", "Semáforo"], 1):
            c = ws.cell(fila, j, t)
            c.font = Font(name=FUENTE, bold=True, size=10, color="FFFFFF")
            c.fill = relleno(NAVY)
            c.alignment = Alignment(horizontal="center")
            c.border = BORDE
        ini_tabla = fila + 1
        mayor = ind.sentido == "mayor"
        for idx, per in enumerate(periodos):
            fila += 1
            m = per["medicion"] if per["medicion"] and per["medicion"].estado != "observado" else None
            mn = ind.meta if mayor else ind.meta_tope
            mx = ind.meta_tope if mayor else ind.meta
            vals = [per["corta"], m.resultado if m else None, mn, mx]
            for j, v in enumerate(vals, 1):
                c = ws.cell(fila, j, v)
                c.border = BORDE
                c.font = Font(name=FUENTE, size=10, bold=(j == 2))
                c.alignment = Alignment(horizontal="center" if j != 1 else "left", indent=1 if j == 1 else 0)
                if j > 1:
                    c.number_format = fmt
            if m:
                # el resultado se recalcula desde los datos de la columna del período; si la fórmula no tiene
                # equivalente en Excel, queda el valor calculado por el sistema
                col = get_column_letter(2 + idx)
                celdas = {v["clave"]: f"{col}{fila_variable[v['clave']]}" for v in ind.variables}
                celdas.update({c["clave"]: repr(c["valor"]) for c in m.constantes_aplicadas})
                formula = a_excel(m.formula_aplicada, celdas, m.precision)
                if formula:
                    ws.cell(fila, 2).value = formula
                if m.meta_tope is None and m.limite_tope is None:
                    ws.cell(fila, 5).value = formula_semaforo(f"B{fila}", m.meta, m.limite, m.sentido == "mayor")
                    ws.cell(fila, 5).alignment = Alignment(horizontal="center", vertical="center")
                else:
                    celda_semaforo(ws.cell(fila, 5), m.semaforo)
            ws.cell(fila, 5).border = BORDE
        if periodos:
            colorear_semaforo(ws, f"E{ini_tabla}:E{fila}")
        fin_tabla = fila
        if periodos:
            ch = LineChart()
            ch.title = ind.nombre[:60]
            ch.height, ch.width = 7.5, 15.5
            ch.add_data(Reference(ws, min_col=2, max_col=4, min_row=ini_tabla - 1, max_row=fin_tabla), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=ini_tabla, max_row=fin_tabla))
            estilos = (("0284C7", None, True), ("16A34A", "dash", False), ("DC2626", "dash", False))
            for sr, (col, dash, marca) in zip(ch.series, estilos, strict=False):
                sr.graphicalProperties.line.solidFill = col
                sr.graphicalProperties.line.width = 28000 if marca else 19000
                if dash:
                    sr.graphicalProperties.line.dashStyle = dash
                sr.smooth = False
                if marca:
                    sr.marker.symbol, sr.marker.size = "circle", 7
                    sr.marker.graphicalProperties.solidFill = col
                else:
                    sr.marker.symbol = "none"
            ch.x_axis.delete = ch.y_axis.delete = False
            ch.legend.position = "b"
            ws.add_chart(ch, f"G{ini_tabla - 1}")
        fila = max(fila, ini_tabla + 13)

        # --- análisis y evaluación
        fila += 1
        _banda(ws, fila, "ANÁLISIS Y EVALUACIÓN DE RESULTADOS")
        fila += 1
        ws.cell(fila, 1, "Período")
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=7)
        ws.cell(fila, 2, "Resultados del análisis")
        ws.merge_cells(start_row=fila, start_column=8, end_row=fila, end_column=13)
        ws.cell(
            fila, 8, "Correcciones, oportunidades o necesidades de mejora, propuestas de cambios y adecuación de recursos"
        )
        for j in range(1, 14):
            c = ws.cell(fila, j)
            c.font = Font(name=FUENTE, bold=True, size=10, color="FFFFFF")
            c.fill = relleno(NAVY)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDE
        ws.row_dimensions[fila].height = 32
        for per in periodos:
            fila += 1
            m = per["medicion"] if per["medicion"] and per["medicion"].estado != "observado" else None
            ws.cell(fila, 1, per["etiqueta"]).font = Font(name=FUENTE, bold=True, size=10)
            ws.cell(fila, 1).alignment = Alignment(vertical="top", indent=1, wrap_text=True)
            ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=7)
            ws.merge_cells(start_row=fila, start_column=8, end_row=fila, end_column=13)
            ws.cell(fila, 2, m.analisis if m else "")
            ws.cell(fila, 8, m.plan_mejora if m else "")
            for j in range(1, 14):
                c = ws.cell(fila, j)
                c.border = BORDE
                if j in (2, 8):
                    c.font = Font(name=FUENTE, size=10)
                    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
            ws.row_dimensions[fila].height = max(
                20, 14 * max(_lineas(m.analisis if m else "", 66), _lineas(m.plan_mejora if m else "", 66))
            )

        # --- notas y pie
        fila += 2
        _banda(ws, fila, "NOTAS")
        fila += 1
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=13)
        notas = ind.notas or ""
        ws.cell(fila, 1, notas)
        ws.cell(fila, 1).alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.cell(fila, 1).font = Font(name=FUENTE, size=10)
        ws.row_dimensions[fila].height = max(30, 14 * _lineas(notas, 150))
        fila += 2
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=13)
        ws.cell(
            fila, 1, f"{params.get('codigo_formato', '')} · Generado el {ahora:%d/%m/%Y %H:%M} por {generado_por} · SIGT"
        )
        ws.cell(fila, 1).font = Font(name=FUENTE, size=8, color=GRIS)
        ws.cell(fila, 1).alignment = Alignment(horizontal="right")
        ws.print_area = f"A1:M{fila}"
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.oddFooter.right.text = f"{params.get('codigo_formato', '')}"
    return wb
