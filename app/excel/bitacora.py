"""Bitácora del sistema en Excel (solo administrador)."""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .comun import (
    BORDE,
    FUENTE,
    GRIS,
    ZEBRA,
    ajustar,
    configurar_impresion,
    encabezado_tabla,
    franja_titulo,
    hora_local,
    relleno,
)


def reporte_bitacora(registros, codigos: dict, generado_por: str, verificacion: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bitácora"
    ahora = hora_local(datetime.utcnow())
    estado = "ÍNTEGRA ✔" if verificacion["ok"] else f"ALTERADA ✖ (registro #{verificacion['primer_error_id']})"
    cols = ["#", "Fecha y hora", "Usuario", "Rol", "Acción", "Entidad", "Indicador", "Detalle", "Hash SHA-256"]
    franja_titulo(
        ws,
        "Bitácora de trazabilidad del sistema",
        f"Generado el {ahora:%d/%m/%Y %H:%M} por {generado_por} · Cadena de hashes: {estado} · "
        f"{verificacion['total']} registros",
        len(cols),
    )
    encabezado_tabla(ws, 4, cols, [8, 17, 28, 11, 22, 12, 11, 90, 24])
    for k, b in enumerate(registros, 5):
        det = "; ".join(f"{a}: {v}" for a, v in (b.detalle or {}).items())
        fila = [
            b.id,
            hora_local(b.fecha),
            b.usuario_nombre,
            b.rol,
            b.accion,
            b.entidad,
            codigos.get(b.indicador_id, ""),
            det,
            b.hash[:16] + "…",
        ]
        for j, v in enumerate(fila, 1):
            c = ws.cell(k, j, v)
            c.font = Font(name="Consolas" if j == 9 else FUENTE, size=10, color=GRIS if j == 9 else "000000")
            c.border = BORDE
            c.alignment = Alignment(vertical="center", wrap_text=(j == 8))
            if k % 2:
                c.fill = relleno(ZEBRA)
        ws.cell(k, 2).number_format = "dd/mm/yyyy hh:mm:ss"
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(5, 4 + len(registros))}"
    configurar_impresion(ws, 4)
    ajustar(ws)
    return wb
