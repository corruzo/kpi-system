"""Integridad de los reportes: fórmula del semáforo, hoja «Control» firmada, protección con contraseña y verificación.

Flujo al exportar:  sellar() → proteger() → a_bytes().   Para comprobar un archivo recibido: verificar().
La firma (HMAC con la clave secreta del sistema) cubre todas las celdas de las demás hojas, fórmulas incluidas.
"""
import hashlib
import io
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font
from openpyxl.workbook.protection import WorkbookProtection

from ..security import firmar_reporte
from .comun import BORDE, FUENTE, GRIS, NAVY, SEM, franja_titulo, relleno

HOJA_CONTROL = "Control"


# ------------------------------------------------------------------ semáforo con fórmula
def formula_semaforo(resultado: str, meta, limite, mayor: bool) -> str:
    """Fórmula de Excel con la misma regla que `dominio.semaforo.clasificar` (sin topes)."""
    op = ">=" if mayor else "<="
    return (f'=IF({resultado}="","",IF({resultado}{op}{meta!r},"Cumple",'
            f'IF({resultado}{op}{limite!r},"Alerta","Crítico")))')


def colorear_semaforo(ws, rango: str) -> None:
    """Color de fondo según el texto de la celda (Cumple / Alerta / Crítico)."""
    for fondo, texto, etiqueta in SEM.values():
        ws.conditional_formatting.add(rango, CellIsRule(
            operator="equal", formula=[f'"{etiqueta}"'], fill=relleno(fondo), font=Font(bold=True, color=texto)))


# ------------------------------------------------------------------ huella del contenido
def _normal(valor) -> str:
    if isinstance(valor, str) and valor.startswith("="):
        return valor.replace(" ", "").upper()          # Excel puede reacomodar espacios y mayúsculas
    if isinstance(valor, datetime):
        return valor.replace(microsecond=0).isoformat()
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)                              # el XML guarda 2.0 como 2
    return repr(valor) if isinstance(valor, float) else str(valor)


def huella(wb) -> str:
    """SHA-256 de todas las celdas con contenido de las hojas de datos (todas menos «Control»)."""
    h = hashlib.sha256()
    for ws in wb.worksheets:
        if ws.title == HOJA_CONTROL:
            continue
        for fila in ws.iter_rows():
            for c in fila:
                if c.value is not None and c.value != "":
                    h.update(f"{ws.title}|{c.coordinate}|{_normal(c.value)}\n".encode())
    return h.hexdigest()


# ------------------------------------------------------------------ sellar, proteger, verificar
def sellar(wb, numero: str, tipo: str, generado_por: str, fecha: datetime) -> str:
    """Agrega la hoja «Control» (número, quién y cuándo, huella y firma) y devuelve la huella."""
    h = huella(wb)
    ws = wb.create_sheet(HOJA_CONTROL)
    franja_titulo(ws, "Control del reporte", "Este archivo fue generado por el SIGT y puede verificarse en el sistema (Reportes → Verificar).", 2)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 78
    filas = [("Número de reporte", numero), ("Tipo", tipo), ("Generado por", generado_por),
             ("Fecha de emisión", f"{fecha:%d/%m/%Y %H:%M}"), ("Huella del contenido", h),
             ("Firma del sistema", firmar_reporte(f"{numero}|{h}"))]
    for k, (etiqueta, valor) in enumerate(filas, 4):
        e, v = ws.cell(k, 1, etiqueta), ws.cell(k, 2, valor)
        e.font, v.font = Font(name=FUENTE, bold=True, color=NAVY), Font(name="Consolas" if k > 7 else FUENTE, size=9 if k > 7 else 10)
        e.border = v.border = BORDE
        v.alignment = Alignment(wrap_text=True, vertical="center")
    ws.cell(11, 1, "Cualquier cambio en las demás hojas invalida la firma.").font = Font(name=FUENTE, italic=True, size=9, color=GRIS)
    return h


def proteger(wb, clave: str) -> None:
    """Bloquea todas las hojas (y la estructura del libro) con la contraseña indicada; filtros y anchos siguen disponibles."""
    for ws in wb.worksheets:
        ws.protection.sheet = True
        ws.protection.set_password(clave)
        ws.protection.autoFilter = False
        ws.protection.sort = False
        ws.protection.formatColumns = False
        ws.protection.formatRows = False
    wb.security = WorkbookProtection(lockStructure=True)
    wb.security.set_workbook_password(clave)


def leer_control(contenido: bytes) -> dict:
    """Datos de la hoja «Control» y huella recalculada; {} si el archivo no es un reporte del SIGT."""
    try:
        wb = load_workbook(io.BytesIO(contenido))
    except Exception:
        return {}
    if HOJA_CONTROL not in wb.sheetnames:
        return {}
    ws = wb[HOJA_CONTROL]
    datos = {str(ws.cell(k, 1).value): ws.cell(k, 2).value for k in range(4, 10)}
    return {"numero": datos.get("Número de reporte"), "tipo": datos.get("Tipo"), "generado_por": datos.get("Generado por"),
            "fecha": datos.get("Fecha de emisión"), "huella_declarada": datos.get("Huella del contenido"),
            "firma": datos.get("Firma del sistema"), "huella_real": huella(wb)}
