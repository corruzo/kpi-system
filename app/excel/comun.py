"""Piezas comunes de los reportes en Excel: paleta, estilos, encabezados, impresión, logo y guardado."""

import io
import math
import textwrap
from datetime import UTC, datetime, timedelta

from openpyxl.drawing.image import Image as XLImagen
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY, BRAND, GRIS, ZEBRA = "0F172A", "0284C7", "64748B", "F8FAFC"
SEM = {
    "verde": ("DCFCE7", "166534", "Cumple"),
    "amarillo": ("FEF3C7", "92400E", "Alerta"),
    "rojo": ("FEE2E2", "991B1B", "Crítico"),
}
ESTADO = {"enviado": "En revisión", "aprobado": "Aprobado", "observado": "Devuelto"}
FUENTE = "Calibri"
_lado = Side(style="thin", color="E2E8F0")
BORDE = Border(left=_lado, right=_lado, top=_lado, bottom=_lado)


def hora_local(dt):
    """UTC → hora de Venezuela (UTC−4, sin horario de verano)."""
    return (dt.replace(tzinfo=UTC) + timedelta(hours=-4)).replace(tzinfo=None) if dt else None


def formato_numero(unidad: str, dec: int) -> str:
    base = "#,##0" + ("." + "0" * dec if dec else "")
    return {"%": base + '"%"', "USD": '"$"' + base, "días": base + '" d"', "TM": base + '" TM"'}.get(
        unidad, base + (f' "{unidad}"' if unidad and unidad != "cantidad" else "")
    )


def relleno(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def encabezado_tabla(ws, fila, textos, ancho=None):
    for j, t in enumerate(textos, 1):
        c = ws.cell(fila, j, t)
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF", size=10)
        c.fill = relleno(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
    ws.row_dimensions[fila].height = 32
    for j, w in enumerate(ancho or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def franja_titulo(ws, titulo, subtitulo, ncols, logo: bytes | None = None):
    """Título y subtítulo de la hoja. Con logo, la primera columna queda para la imagen (fondo blanco) y el
    texto empieza en la segunda, de modo que nunca se superponen."""
    ws.sheet_view.showGridLines = False
    primera = 1
    if logo and ncols > 1:
        primera = 2
        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
        ws.cell(1, 1).fill = relleno("FFFFFF")
        ws.cell(1, 1).border = BORDE
        poner_logo(ws, logo, "A1", 40, max_ancho=170)
    ws.merge_cells(start_row=1, start_column=primera, end_row=1, end_column=ncols)
    c = ws.cell(1, primera, titulo)
    c.font = Font(name=FUENTE, bold=True, size=16, color="FFFFFF")
    c.fill = relleno(NAVY)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=primera, end_row=2, end_column=ncols)
    c = ws.cell(2, primera, subtitulo)
    c.font = Font(name=FUENTE, size=10, color="FFFFFF")
    c.fill = relleno(BRAND)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 20


def configurar_impresion(ws, fila_encabezado, apaisado=True):
    ws.page_setup.orientation = "landscape" if apaisado else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{fila_encabezado}:{fila_encabezado}"
    ws.oddFooter.center.text = "Tripoliven, C.A. · SIGT · Página &P de &N"
    ws.page_margins.left = ws.page_margins.right = 0.4


def celda_semaforo(c, semaforo):
    bg, fg, txt = SEM[semaforo]
    c.value = txt
    c.fill = relleno(bg)
    c.font = Font(name=FUENTE, bold=True, color=fg, size=10)
    c.alignment = Alignment(horizontal="center", vertical="center")


def poner_logo(ws, logo: bytes | None, celda: str = "A1", alto: int = 36, max_ancho: int = 0) -> int:
    """Inserta el logo de la empresa (PNG) si existe; devuelve su ancho en píxeles (0 si no hay logo)."""
    if not logo:
        return 0
    try:
        img = XLImagen(io.BytesIO(logo))
        img.width, img.height = int(img.width * alto / img.height), alto
        if max_ancho and img.width > max_ancho:           # logos muy anchos se reducen para caber en su espacio
            img.height, img.width = int(img.height * max_ancho / img.width), max_ancho
        ws.add_image(img, celda)
        return img.width
    except Exception:
        return 0


def ubicar_grafico(ch, ws, pos_fila: int, ancho_cm: float, alto_cm: float, columna: str = "A") -> int:
    """Coloca un gráfico con la leyenda debajo y sin invadir el área de trazado; devuelve cuántas filas ocupa
    (las filas miden 15 pt = 0,53 cm), para que lo siguiente empiece después de él."""
    from openpyxl.chart.layout import Layout, ManualLayout
    ch.width, ch.height = ancho_cm, alto_cm
    if ch.legend is not None:
        ch.legend.position = "b"
        ch.legend.overlay = False
    if ch.title is not None:
        ch.title.overlay = False
    ch.plot_area.layout = Layout(manualLayout=ManualLayout(layoutTarget="inner", xMode="edge", yMode="edge",
                                                           x=0.09, y=0.16, w=0.88, h=0.58))
    ws.add_chart(ch, f"{columna}{pos_fila}")
    return math.ceil(alto_cm / 0.53) + 1


# ------------------------------------------------------------------ ajuste de filas y columnas al contenido
_PT_POR_LINEA = 1.32     # alto de una línea de texto, en múltiplos del tamaño de fuente
_ANCHO_MIN, _ANCHO_MAX = 8, 60


def _texto(c) -> str:
    v = c.value
    if v is None:
        return ""
    if isinstance(v, datetime):
        return "00/00/0000 00:00"
    if isinstance(v, float):
        return f"{v:,.2f}"
    if isinstance(v, str) and v.startswith("="):
        return "0000.00"                       # lo que muestra una fórmula: un número
    return str(v)


def _caracteres_por_linea(ancho: float, c) -> int:
    tam = (c.font.sz or 11) if c.font else 11
    mono = bool(c.font and c.font.name and c.font.name.lower() in ("consolas", "courier new"))
    por_unidad = (0.88 if mono else 1.05) * 11 / tam
    sangria = int((c.alignment.indent or 0) * 2) if c.alignment else 0
    return max(4, int(ancho * por_unidad) - sangria)


def lineas_necesarias(c, ancho: float) -> int:
    """Líneas que ocupa el texto de la celda en un ancho dado (en unidades de columna de Excel)."""
    texto = _texto(c)
    if not texto:
        return 1
    cpl = _caracteres_por_linea(ancho, c)
    return sum(max(1, len(textwrap.wrap(p, cpl, break_long_words=True)) if p else 1) for p in texto.split("\n"))


def altura_necesaria(c, ancho: float) -> float:
    tam = (c.font.sz or 11) if c.font else 11
    return lineas_necesarias(c, ancho) * tam * _PT_POR_LINEA + 4


def _ancho_columna(ws, letra: str) -> float:
    return ws.column_dimensions[letra].width or 8.43


def ajustar(ws, desde_fila: int = 1, columnas: bool = True) -> None:
    """Deja cada columna y cada fila del tamaño que pide su contenido, sin recortes ni desbordes.

    - Columnas (si `columnas`): crecen hasta el texto más largo de las celdas sin ajuste de línea (tope 60).
    - Filas: las celdas con ajuste de línea o saltos de línea fijan el alto de su fila; las combinadas usan la
      suma de los anchos que abarcan, y si abarcan varias filas el alto se reparte entre ellas.
    """
    combinadas = {}
    for rango in ws.merged_cells.ranges:
        combinadas[(rango.min_row, rango.min_col)] = rango
    cubiertas = {(r, c) for rg in ws.merged_cells.ranges for r in range(rg.min_row, rg.max_row + 1)
                 for c in range(rg.min_col, rg.max_col + 1)} - set(combinadas)

    if columnas:
        for col in range(1, ws.max_column + 1):
            letra = get_column_letter(col)
            maximo = 0.0
            for fila in range(desde_fila, ws.max_row + 1):
                c = ws.cell(fila, col)
                if (fila, col) in cubiertas or (fila, col) in combinadas or c.value is None:
                    continue
                if sum(1 for k in range(1, ws.max_column + 1) if ws.cell(fila, k).value is not None) == 1:
                    continue                      # título de sección: se extiende sobre las celdas vacías contiguas
                if c.alignment and c.alignment.wrap_text:
                    palabra = max((len(p) for p in _texto(c).split()), default=0)     # al menos la palabra más larga
                    maximo = max(maximo, min(palabra, 30) * 1.1)
                    continue
                negrita = 1.1 if (c.font and c.font.b) else 1.0
                maximo = max(maximo, max(len(t) for t in _texto(c).split("\n")) * 1.05 * negrita * (((c.font.sz or 11) if c.font else 11) / 11) + 2.5)
            actual = ws.column_dimensions[letra].width or 0
            ws.column_dimensions[letra].width = round(min(_ANCHO_MAX, max(actual, _ANCHO_MIN if maximo else 0, maximo)), 1) or None

    necesidad: dict[int, float] = {}
    for fila in range(desde_fila, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            c = ws.cell(fila, col)
            if (fila, col) in cubiertas or c.value is None or not (c.alignment and c.alignment.wrap_text):
                continue
            rg = combinadas.get((fila, col))
            ancho = sum(_ancho_columna(ws, get_column_letter(k)) for k in range(rg.min_col, rg.max_col + 1)) if rg else \
                _ancho_columna(ws, get_column_letter(col))
            alto = altura_necesaria(c, ancho)
            filas = range(fila, (rg.max_row if rg else fila) + 1)
            for f in filas:                                  # el alto de una combinada se reparte entre sus filas
                necesidad[f] = max(necesidad.get(f, 0), alto / len(filas))
    for fila, alto in necesidad.items():
        ws.row_dimensions[fila].height = round(max(ws.row_dimensions[fila].height or 15, alto), 1)


def a_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ reporte general
