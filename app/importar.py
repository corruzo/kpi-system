"""Importación de las fichas históricas en Excel («Ficha técnica de indicador de gestión», una hoja por indicador).

Lee la estructura de la plantilla del SGC (definición, semáforo, responsables, resultados por período y análisis),
traduce la fórmula de la celda de cálculo a la sintaxis del sistema y propone indicadores y mediciones.
Nada se guarda aquí: el resultado es una vista previa que el administrador revisa y confirma.
"""
import io
import re
import unicodedata
from datetime import date

from openpyxl import load_workbook

from . import calc

MAX_BYTES = 8 * 1024 * 1024
SEGUIMIENTO = {"desempeño de proceso": "proceso", "objetivo organizacional": "objetivo",
               "desempeño de proveedor externo": "proveedor"}
UNIDADES = {"dia": "días", "día": "días", "dias": "días", "días": "días", "cantidad": "cantidad",
            "porcentaje": "%", "%": "%"}
MESES = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8, "sep": 9, "set": 9,
         "oct": 10, "nov": 11, "dic": 12}
FUNCS = {"IF": "si", "ROUND": "round", "SUM": "sum", "AVERAGE": "avg", "MIN": "min", "MAX": "max",
         "ABS": "abs", "SQRT": "sqrt"}


def _txt(v) -> str:
    return "" if v is None else re.sub(r"[ \t\r\f\v]+", " ", str(v)).strip()


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip().replace(",", "."))
    except ValueError:
        return None


def _norm(t: str) -> str:
    t = unicodedata.normalize("NFD", t.lower())
    return re.sub(r"[^a-z0-9 ]+", " ", "".join(c for c in t if unicodedata.category(c) != "Mn")).strip()


def _slug(t: str) -> str:
    return re.sub(r"\s+", "_", _norm(t))[:30].strip("_") or "valor"


def _suma_meses(d: date, n: int) -> date:
    y, m = divmod(d.month - 1 + n, 12)
    return date(d.year + y, m + 1, 1)


# ------------------------------------------------------------------ utilidades de hoja
def _filas_por_etiqueta(ws, col=2, hasta=120) -> dict[str, list[int]]:
    out: dict[str, list[int]] = {}
    for r in range(1, hasta):
        t = _norm(_txt(ws.cell(r, col).value))
        if t:
            out.setdefault(t, []).append(r)
    return out


def _lista_columna(ws, col: int, etiqueta: str, maximo=200) -> list[str]:
    """Valores no vacíos bajo una cabecera (ej. «Gerencias:») hasta la primera celda vacía."""
    for r in range(1, 400):
        if _norm(_txt(ws.cell(r, col).value)) == _norm(etiqueta):
            out = []
            for k in range(r + 1, r + maximo):
                t = _txt(ws.cell(k, col).value)
                if not t:
                    break
                out.append(t)
            return out
    return []


def _catalogos(ws) -> dict:
    cat = {"gerencia": _lista_columna(ws, 13, "Gerencias:"), "area": _lista_columna(ws, 20, "Departamentos:"),
           "cargo": _lista_columna(ws, 25, "Cargos:"), "objetivo": [], "proceso": []}
    for r in range(1, 200):  # objetivos organizacionales: columnas B (N°) y C (texto)
        if _txt(ws.cell(r, 2).value) == "N°" and _norm(_txt(ws.cell(r, 3).value)) == "objetivo":
            k = r + 1
            while _txt(ws.cell(k, 3).value):
                cat["objetivo"].append(_txt(ws.cell(k, 3).value))
                k += 1
            break
    for r in range(200, 400):  # procesos: lista en la columna F, tras una línea de marcas «xxxx»
        if _txt(ws.cell(r, 1).value).lower().startswith("xxxx"):
            k = r + 2
            while _txt(ws.cell(k, 6).value):
                cat["proceso"].append(_txt(ws.cell(k, 6).value))
                k += 1
            break
    return cat


# ------------------------------------------------------------------ fórmula
def _traducir_formula(f: str, mapa: dict[int, str]) -> str:
    """Convierte la fórmula de la celda de cálculo (p. ej. =IF(D23=0," ",D22*100/D23)) a la del sistema."""
    f = (f or "").strip().lstrip("=").strip()
    m = re.match(r'^IF\(\s*[^,]+?\s*=\s*0\s*,\s*"\s*"\s*,\s*(.+)\)$', f, re.I | re.S)
    if m:  # protección contra denominador 0: el sistema ya impide enviar datos inválidos
        f = m.group(1)

    def celda(mm):
        fila = int(mm.group(2))
        if fila not in mapa:
            raise ValueError(f"la celda {mm.group(0)} no corresponde a una variable de la ficha")
        return mapa[fila]
    f = re.sub(r"\$?\b([A-Z]{1,2})\$?(\d{1,3})\b(?!\()", celda, f)
    f = re.sub(r"\b([A-Za-z]+)\(", lambda mm: (FUNCS[mm.group(1).upper()] if mm.group(1).upper() in FUNCS
                                               else (_ for _ in ()).throw(ValueError(f"función «{mm.group(1)}»"))) + "(", f)
    f = f.replace(";", ",").replace("<>", "!=")
    f = re.sub(r"(?<![<>!=])=(?!=)", "==", f)
    return calc.normalizar(f)


def _simbolos(doc: str) -> dict[str, str]:
    """«SG = solicitudes gestionadas» → {'sg': 'solicitudes gestionadas'} (leyenda de la fórmula)."""
    out = {}
    for linea in (doc or "").splitlines():
        m = re.match(r"^\s*([A-Za-zÁÉÍÓÚáéíóú]{1,8})\s*=\s*(.+)$", linea)
        if m:
            out[_norm(m.group(1)).replace(" ", "")] = _norm(m.group(2))
    return out


def _claves(etiquetas: list[str], doc: str, sigla: str) -> list[str]:
    """Clave de cada variable: el símbolo de la leyenda si coincide con la etiqueta; si no, un nombre derivado."""
    sim, usadas, out = _simbolos(doc), set(), []
    for et in etiquetas:
        e = set(_norm(et).split())
        mejor, puntaje = None, 0.0
        for s, desc in sim.items():
            d = set(desc.split())
            j = max(len(e & d) / max(1, len(e | d)), len(e & d) / max(1, len(e)))  # similitud o contención
            if j > puntaje and s != _norm(sigla).replace(" ", "") and s not in usadas:
                mejor, puntaje = s, j
        k = mejor if mejor and puntaje >= 0.6 and re.match(r"^[a-z][a-z0-9]{1,9}$", mejor) else _slug(et)
        if not re.match(r"^[a-z]", k):
            k = "v_" + k
        base, n = k, 2
        while k in usadas or k in calc.RESERVADAS:
            k, n = f"{base}_{n}", n + 1
        usadas.add(k)
        out.append(k)
    return out


# ------------------------------------------------------------------ hoja → indicador
def _hoja(ws, wv, nombre_hoja: str) -> dict | None:
    if not any("ficha t" in _norm(_txt(ws.cell(r, c).value)) for r in range(1, 6) for c in range(1, 15)):
        return None
    av: list[str] = []
    f = _filas_por_etiqueta(ws)
    try:
        r_nom = f["nombre del indicador"][0]
        r_res = f["resultados de las determinaciones"][0]
        r_val = f["valores del indicador"][0]
    except KeyError:
        return None
    V = lambda r, c: wv.cell(r, c).value  # noqa: E731 (valores calculados)

    # seguimiento (celda bajo «Seguimiento de:»)
    seg = "proceso"
    for r in range(1, 6):
        for c in range(1, 30):
            if _norm(_txt(ws.cell(r, c).value)).startswith("seguimiento de"):
                seg = SEGUIMIENTO.get(_txt(V(r + 1, c)).lower(), "proceso")
    nombre = _txt(V(r_nom, 4))
    sigla = (re.search(r"\(([A-Za-z0-9ÁÉÍÓÚ]{1,10})\)\s*$", nombre) or [None, ""])[1]
    r_u = r_nom + 1
    unidad_raw = _txt(V(r_u, 4))
    unidad = UNIDADES.get(unidad_raw.lower(), unidad_raw[:12] or "cantidad")
    meta_txt = _txt(V(r_u, 11))
    m = re.search(r"(-?\d+(?:[.,]\d+)?)\s*(m[ií]n|m[aá]x)", meta_txt, re.I)
    if m:
        meta, sentido = float(m.group(1).replace(",", ".")), "mayor" if m.group(2).lower().startswith("mí") or \
            m.group(2).lower().startswith("mi") else "menor"
    else:
        meta, sentido = _num(V(r_u, 11)) or 0.0, "mayor"
        av.append(f"No se pudo interpretar la meta «{meta_txt}»: verifique meta y sentido.")
    frec = _txt(V(r_u, 16)).lower()
    if frec not in calc.FRECUENCIAS:
        av.append(f"Frecuencia «{frec}» no reconocida: se usó «mensual».")
        frec = "mensual"
    # semáforo (mín./máx. por franja)
    vv = lambda r, c: _num(V(r_nom + r, c))  # noqa: E731
    v13, w13, v14, w14 = vv(2, 22), vv(2, 23), vv(3, 22), vv(3, 23)
    if sentido == "mayor":
        meta_tope = w13 if w13 is not None and w13 >= meta else None
        limite = v14 if v14 is not None and v14 <= meta else meta
        limite_tope = w14 if w14 is not None and w14 >= limite else None
    else:
        meta_tope = v13 if v13 is not None and v13 <= meta else None
        limite = w14 if w14 is not None and w14 >= meta else meta
        limite_tope = v14 if v14 is not None and v14 <= limite else None
    try:
        calc.validar_umbrales(meta, limite, sentido, meta_tope, limite_tope)
    except ValueError:
        meta_tope = limite_tope = None
    r_f = (f.get("formula de calculo") or [r_nom + 2])[0]
    doc = _txt(ws.cell(r_f, 4).value) if ws.cell(r_f, 4).value else _txt(V(r_f, 4))
    doc_raw = str(V(r_f, 4) or "")
    r_fu = (f.get("fuente de la informacion") or [r_nom + 5])[0]
    fuente, ejercicio_txt = _txt(V(r_fu, 4)), _txt(V(r_fu, 22))
    r_g = (f.get("gerencia") or [r_val - 6])[0]
    gerencia, area, lider = _txt(V(r_g, 4)), _txt(V(r_g, 13)), _txt(V(r_g, 20))
    resp_med, resp_reg, resp_ana = _txt(V(r_g + 1, 4)), _txt(V(r_g + 1, 13)), _txt(V(r_g + 1, 20))
    # proceso/objetivo y su descripción: las dos celdas combinadas bajo la cabecera de la ficha (columna B)
    etiquetas_b = [r for r in range(1, r_nom) if _txt(V(r, 2)) and _norm(_txt(V(r, 2))) in (
        "proceso", "objetivo del proceso", "objetivo organizacional", "objetivo especifico", "bien o servicio")]
    ref_nombre = _txt(V(etiquetas_b[0] + 1, 2)) if etiquetas_b else _txt(V(7, 2))
    ref_detalle = _txt(V(etiquetas_b[1] + 1, 2)) if len(etiquetas_b) > 1 else _txt(V(9, 2))

    # períodos (columnas con encabezado en la fila siguiente a «Resultados de las determinaciones»)
    r_h = r_res + 1
    slots = [(c, _txt(V(r_h, c))) for c in range(4, 40) if _txt(V(r_h, c))][:12]
    validos = [(c, e) for c, e in slots if _norm(e) != "no aplica"]
    if not validos:
        av.append("La hoja no tiene períodos definidos.")
    # variables
    filas_var = [r for r in range(r_res + 2, r_val) if _txt(ws.cell(r, 2).value)]
    etiquetas = [_txt(ws.cell(r, 2).value) for r in filas_var]
    claves = _claves(etiquetas, doc_raw, sigla)
    mapa = dict(zip(filas_var, claves))
    # fórmula
    formula, celda_f = "", ws.cell(r_val + 2, 3).value
    try:
        formula = _traducir_formula(str(celda_f), mapa)
        calc.validar_formula(formula, claves)
    except (ValueError, calc.FormulaError) as e:
        av.append(f"No se pudo traducir la fórmula «{celda_f}» ({e}). Escríbala en el recuadro de fórmula.")
        formula = ""
    # calendario del ejercicio
    k = calc.MESES_POR_PERIODO[frec]
    inicio = None
    y0 = None
    if validos:
        mes_fin = MESES.get(_norm(validos[0][1])[:3])
        if mes_fin:
            inicio = (mes_fin - k) % 12 + 1
    mm = re.search(r"(\d{4})\s*[-–/]\s*(\d{4})", ejercicio_txt)
    if mm:
        y0, y1 = int(mm.group(1)), int(mm.group(2))
    else:
        mm = re.search(r"(\d{4})", ejercicio_txt)
        y1 = int(mm.group(1)) if mm else None
        y0 = y1 - (0 if inicio == 1 else 1) if y1 and inicio else None
    if y0 is None or inicio is None:
        av.append("No se pudo determinar el ejercicio fiscal: no se importarán mediciones.")
    # análisis
    r_an = (f.get("analisis y evaluacion de resultados") or [r_val + 12])[0]
    r_ah = next((r for r in range(r_an + 1, r_an + 8) if _norm(_txt(ws.cell(r, 2).value)) == "periodo"), r_an + 3)
    # mediciones
    mediciones, dec = [], 2
    if formula and inicio and y0:
        for i, (c, etq) in enumerate(validos):
            vals, ok = {}, True
            for r, clave in zip(filas_var, claves):
                x = _num(V(r, c))
                if x is None:
                    ok = False
                    break
                vals[clave] = x
            if not ok:
                continue
            fin = _suma_meses(date(y0, inicio, 1), (i + 1) * k - 1)  # mes de cierre del período
            if frec in ("mensual", "eventual"):
                per = f"{fin.year}-{fin.month:02d}"
            else:
                y_fin = y0 + (0 if inicio == 1 else 1)
                per = str(y_fin) if frec == "anual" else f"{y_fin}-{calc.PREFIJO[frec]}{i + 1}"
            try:
                res = calc.evaluar(formula, vals, {}, dec)
            except calc.FormulaError as e:
                av.append(f"{etq}: no se importó ({e})")
                continue
            ra = r_ah + 1 + slots.index((c, etq))
            mediciones.append({"periodo": per, "etiqueta": etq, "valores": vals, "resultado": res,
                               "semaforo": calc.semaforo(res, meta, limite, sentido, meta_tope, limite_tope),
                               "analisis": _txt(V(ra, 4)), "plan_mejora": _txt(V(ra, 13))})
    notas = ""
    r_nt = (f.get("notas") or [0])[0]
    if r_nt:
        extras = [_txt(V(r, c)) for r in range(r_nt + 1, r_nt + 12) for c in range(2, 24)]
        notas = " ".join(t for t in extras if t and not re.match(r"^NC-\d+-\d+$", t))
    doc_nota = f"Fórmula documentada en la ficha original: {doc}" if doc else ""
    return {
        "hoja": nombre_hoja, "codigo": "", "nombre": nombre, "sigla": sigla, "area": area, "gerencia": gerencia,
        "seguimiento": seg, "ref_nombre": ref_nombre, "ref_detalle": ref_detalle, "lider": lider,
        "resp_medicion": resp_med, "resp_registro": resp_reg, "resp_analisis": resp_ana,
        "unidad": unidad, "frecuencia": frec, "sentido": sentido, "meta": meta, "limite": limite,
        "meta_tope": meta_tope, "limite_tope": limite_tope, "decimales": dec, "fuente": fuente,
        "objetivo": ref_detalle, "notas": " · ".join(x for x in (notas, doc_nota) if x)[:2000],
        "formula": formula, "variables": [{"clave": c, "etiqueta": e} for c, e in zip(claves, etiquetas)],
        "ejercicio": ejercicio_txt, "inicio_ejercicio": inicio, "mediciones": mediciones, "avisos": av,
    }


def leer_libro(contenido: bytes, archivo: str = "") -> dict:
    if len(contenido) > MAX_BYTES:
        raise ValueError("El archivo supera 8 MB.")
    try:
        wb = load_workbook(io.BytesIO(contenido), data_only=False)
        wv = load_workbook(io.BytesIO(contenido), data_only=True)
    except Exception:
        raise ValueError("No se pudo leer el archivo: debe ser un libro de Excel (.xlsx).")
    inds, avisos = [], []
    for ws in wb.worksheets:
        try:
            r = _hoja(ws, wv[ws.title], ws.title)
        except Exception as e:  # una hoja rara no debe impedir importar las demás
            avisos.append(f"Hoja «{ws.title}»: no se pudo interpretar ({type(e).__name__}).")
            continue
        if r:
            inds.append(r)
        else:
            avisos.append(f"Hoja «{ws.title}»: no tiene el formato de ficha técnica; se omitió.")
    cat = {}
    for ws in wb.worksheets:
        cat = _catalogos(ws)
        if any(cat.values()):
            break
    # el objetivo de cada proceso se toma de las propias fichas
    detalle = {i["ref_nombre"]: i["ref_detalle"] for i in inds if i["seguimiento"] != "objetivo" and i["ref_nombre"]}
    inicios = {i["inicio_ejercicio"] for i in inds if i["inicio_ejercicio"]}
    return {"archivo": archivo, "indicadores": inds, "avisos": avisos, "catalogos": cat,
            "detalle_procesos": detalle, "inicio_ejercicio": inicios.pop() if len(inicios) == 1 else None}
