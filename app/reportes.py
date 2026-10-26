"""Reportes en Excel (.xlsx) con formato corporativo."""
import io
import re
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import calc

NAVY, BRAND, GRIS, ZEBRA = "0F172A", "0284C7", "64748B", "F8FAFC"
SEM = {"verde": ("DCFCE7", "166534", "Cumple"), "amarillo": ("FEF3C7", "92400E", "Alerta"),
       "rojo": ("FEE2E2", "991B1B", "Crítico")}
ESTADO = {"enviado": "En revisión", "aprobado": "Aprobado", "observado": "Devuelto"}
FUENTE = "Calibri"
_lado = Side(style="thin", color="E2E8F0")
BORDE = Border(left=_lado, right=_lado, top=_lado, bottom=_lado)


def _local(dt):
    """UTC → hora de Venezuela (UTC−4, sin horario de verano)."""
    return (dt.replace(tzinfo=timezone.utc) + timedelta(hours=-4)).replace(tzinfo=None) if dt else None


def _fmt_num(unidad: str, dec: int) -> str:
    base = "#,##0" + ("." + "0" * dec if dec else "")
    return {"%": base + '"%"', "USD": '"$"' + base, "días": base + '" d"', "TM": base + '" TM"'}.get(
        unidad, base + (f' "{unidad}"' if unidad and unidad != "cantidad" else ""))


def _fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def _encabezado(ws, fila, textos, ancho=None):
    for j, t in enumerate(textos, 1):
        c = ws.cell(fila, j, t)
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF", size=10)
        c.fill = _fill(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
    ws.row_dimensions[fila].height = 32
    for j, w in enumerate(ancho or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _titulo(ws, titulo, subtitulo, ncols):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(1, 1, titulo)
    c.font = Font(name=FUENTE, bold=True, size=16, color="FFFFFF")
    c.fill = _fill(NAVY)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(2, 1, subtitulo)
    c.font = Font(name=FUENTE, size=10, color="FFFFFF")
    c.fill = _fill(BRAND)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 20


def _impresion(ws, fila_encabezado, apaisado=True):
    ws.page_setup.orientation = "landscape" if apaisado else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{fila_encabezado}:{fila_encabezado}"
    ws.oddFooter.center.text = "Tripoliven, C.A. · SIGT · Página &P de &N"
    ws.page_margins.left = ws.page_margins.right = 0.4


def _celda_sem(c, semaforo):
    bg, fg, txt = SEM[semaforo]
    c.value = txt
    c.fill = _fill(bg)
    c.font = Font(name=FUENTE, bold=True, color=fg, size=10)
    c.alignment = Alignment(horizontal="center", vertical="center")


def _legible(formula, variables):
    t = formula
    for v in sorted(variables, key=lambda v: -len(v["clave"])):
        t = re.sub(r"\b" + re.escape(v["clave"]) + r"\b", v["etiqueta"], t)
    return t


def _guardar(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ reporte general
def reporte_general(mediciones, indicadores, generado_por: str, filtros: str, por_aprobar: int) -> bytes:
    """mediciones: [Medicion]; indicadores: [Indicador] (los incluidos en el reporte)."""
    wb = Workbook()
    ahora = _local(datetime.utcnow())
    sub = f"Reporte de indicadores de gestión · Generado el {ahora:%d/%m/%Y %H:%M} por {generado_por}"

    # ---------- Resumen
    ws = wb.active
    ws.title = "Resumen"
    _titulo(ws, "Tripoliven, C.A. — Sistema de Indicadores de Gestión", sub, 7)
    ws["A4"], ws["B4"] = "Filtros aplicados", filtros or "Ninguno (todos los datos)"
    ws["A4"].font = Font(name=FUENTE, bold=True, color=GRIS)
    ultimo = {}
    for m in sorted(mediciones, key=lambda m: (m.indicador_id, m.periodo)):
        if m.estado in ("enviado", "aprobado"):
            ultimo[m.indicador_id] = m
    conteo = {"verde": 0, "amarillo": 0, "rojo": 0}
    por_area: dict[str, dict] = {}
    for i in indicadores:
        a = por_area.setdefault(i.area.nombre, {"total": 0, "verde": 0, "amarillo": 0, "rojo": 0, "sin": 0})
        a["total"] += 1
        m = ultimo.get(i.id)
        if m:
            conteo[m.semaforo] += 1
            a[m.semaforo] += 1
        else:
            a["sin"] += 1
    total = len(indicadores)
    kpis = [("Indicadores", total), ("Cumplimiento (verde)", conteo["verde"] / total if total else 0),
            ("En alerta", conteo["amarillo"]), ("Críticos", conteo["rojo"]),
            ("Cargas registradas", len(mediciones)), ("Por aprobar", por_aprobar)]
    colores = [NAVY, "16A34A", "F59E0B", "DC2626", BRAND, GRIS]
    for k, ((t, v), col) in enumerate(zip(kpis, colores)):
        r, c = 6 + (k // 3) * 3, 1 + (k % 3) * 2
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 1)
        ws.merge_cells(start_row=r + 1, start_column=c, end_row=r + 1, end_column=c + 1)
        a = ws.cell(r, c, t.upper())
        a.font = Font(name=FUENTE, size=9, bold=True, color="FFFFFF")
        a.fill = _fill(col)
        a.alignment = Alignment(horizontal="center")
        ws.cell(r, c + 1).fill = _fill(col)
        b = ws.cell(r + 1, c, v)
        b.font = Font(name=FUENTE, size=20, bold=True, color=col)
        b.alignment = Alignment(horizontal="center", vertical="center")
        b.fill = _fill("FFFFFF")
        b.border = BORDE
        ws.cell(r + 1, c + 1).border = BORDE
        if t.startswith("Cumpl"):
            b.number_format = "0.0%"
        ws.row_dimensions[r + 1].height = 34
    f0 = 13
    ws.cell(f0 - 1, 1, "Cumplimiento por departamento (última medición de cada indicador)").font = \
        Font(name=FUENTE, bold=True, size=12, color=NAVY)
    _encabezado(ws, f0, ["Área", "Indicadores", "Cumple", "Alerta", "Crítico", "Sin datos", "% cumplimiento"],
                [30, 16, 14, 14, 14, 14, 20])
    for k, (nombre, a) in enumerate(sorted(por_area.items()), f0 + 1):
        fila = [nombre, a["total"], a["verde"], a["amarillo"], a["rojo"], a["sin"],
                a["verde"] / a["total"] if a["total"] else 0]
        for j, v in enumerate(fila, 1):
            c = ws.cell(k, j, v)
            c.font = Font(name=FUENTE, size=10, bold=(j == 1))
            c.border = BORDE
            c.alignment = Alignment(horizontal="left" if j == 1 else "center")
            if k % 2:
                c.fill = _fill(ZEBRA)
        ws.cell(k, 7).number_format = "0.0%"
    fin = f0 + len(por_area)
    if por_area:
        ws.conditional_formatting.add(f"G{f0 + 1}:G{fin}", DataBarRule(start_type="num", start_value=0,
                                                                      end_type="num", end_value=1, color="16A34A"))
        ch = BarChart()
        ch.type, ch.grouping, ch.overlap = "col", "stacked", 100
        ch.title = "Semáforos por departamento"
        ch.add_data(Reference(ws, min_col=3, max_col=5, min_row=f0, max_row=fin), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=f0 + 1, max_row=fin))
        for s, col in zip(ch.series, ("16A34A", "F59E0B", "DC2626")):
            s.graphicalProperties.solidFill = col
            s.graphicalProperties.line.solidFill = col
        ch.x_axis.delete = ch.y_axis.delete = False  # Excel oculta los ejes si no se indica
        ch.y_axis.numFmt, ch.y_axis.majorGridlines = "0", None
        ch.height, ch.width = 8.5, 20
        ws.add_chart(ch, f"A{fin + 3}")
    _impresion(ws, f0, apaisado=False)

    # ---------- Mediciones
    ws = wb.create_sheet("Mediciones")
    cols = ["Código", "Indicador", "Área", "Período", "Frecuencia", "Resultado", "Meta", "Límite crítico", "Sentido",
            "Semáforo", "Estado", "Cargado por", "Fecha de carga", "Revisado por", "Fecha de revisión",
            "Nota de revisión", "Análisis de resultados", "Correcciones / plan de mejora", "Versión"]
    anchos = [10, 38, 20, 16, 12, 14, 12, 14, 12, 12, 13, 24, 17, 24, 17, 34, 44, 44, 9]
    _titulo(ws, "Mediciones registradas", sub, len(cols))
    _encabezado(ws, 4, cols, anchos)
    ordenadas = sorted(mediciones, key=lambda m: (m.indicador.codigo, m.periodo))
    for k, m in enumerate(ordenadas, 5):
        i = m.indicador
        dec = m.decimales if m.decimales is not None else (i.decimales if i.decimales is not None else 2)
        fila = [i.codigo, i.nombre, i.area.nombre, calc.etiqueta(i.frecuencia, m.periodo), i.frecuencia.capitalize(),
                m.resultado, m.meta, m.limite, "Mayor es mejor" if m.sentido == "mayor" else "Menor es mejor", None,
                ESTADO[m.estado], m.cargado_por.nombre, _local(m.cargado_en),
                m.revisado_por.nombre if m.revisado_por else "", _local(m.revisado_en), m.nota_revision,
                m.analisis, m.plan_mejora, m.version]
        for j, v in enumerate(fila, 1):
            c = ws.cell(k, j, v)
            c.font = Font(name=FUENTE, size=10)
            c.border = BORDE
            c.alignment = Alignment(vertical="top", wrap_text=j in (2, 16, 17, 18))
            if k % 2:
                c.fill = _fill(ZEBRA)
        for j in (6, 7, 8):
            ws.cell(k, j).number_format = _fmt_num(i.unidad, dec)
        ws.cell(k, 6).font = Font(name=FUENTE, size=10, bold=True)
        for j in (13, 15):
            ws.cell(k, j).number_format = "dd/mm/yyyy hh:mm"
        _celda_sem(ws.cell(k, 10), m.semaforo)
        ws.cell(k, 10).border = BORDE
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(5, 4 + len(ordenadas))}"
    _impresion(ws, 4)

    # ---------- Tendencias (matriz indicador × período, por frecuencia)
    ws = wb.create_sheet("Tendencias")
    _titulo(ws, "Tendencias por indicador", sub, 12)
    fila = 4
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 38
    for freq in ("mensual", "trimestral", "semestral", "anual"):
        inds = [i for i in indicadores if i.frecuencia == freq]
        if not inds:
            continue
        validas = [m for m in mediciones if m.indicador.frecuencia == freq and m.estado in ("enviado", "aprobado")]
        periodos = sorted({m.periodo for m in validas})
        ws.cell(fila, 1, f"Indicadores {freq}es" if freq != "anual" else "Indicadores anuales").font = \
            Font(name=FUENTE, bold=True, size=12, color=NAVY)
        fila += 1
        _encabezado(ws, fila, ["Código", "Indicador"] + [calc.etiqueta(freq, p) for p in periodos] + ["Meta"])
        for j in range(3, 3 + len(periodos) + 1):
            ws.column_dimensions[get_column_letter(j)].width = max(ws.column_dimensions[get_column_letter(j)].width or 0, 15)
        por = {(m.indicador_id, m.periodo): m for m in validas}
        for i in sorted(inds, key=lambda x: x.codigo):
            fila += 1
            dec = i.decimales if i.decimales is not None else 2
            ws.cell(fila, 1, i.codigo).font = Font(name=FUENTE, bold=True, size=10)
            ws.cell(fila, 2, i.nombre).alignment = Alignment(wrap_text=True, vertical="center")
            for j, p in enumerate(periodos, 3):
                m = por.get((i.id, p))
                c = ws.cell(fila, j, m.resultado if m else None)
                c.number_format = _fmt_num(i.unidad, dec)
                c.alignment = Alignment(horizontal="center", vertical="center")
                if m:
                    bg, fg, _ = SEM[m.semaforo]
                    c.fill = _fill(bg)
                    c.font = Font(name=FUENTE, bold=True, color=fg, size=10)
            c = ws.cell(fila, 3 + len(periodos), i.meta)
            c.number_format = _fmt_num(i.unidad, dec)
            c.font = Font(name=FUENTE, italic=True, color=GRIS, size=10)
            c.alignment = Alignment(horizontal="center", vertical="center")
            for j in range(1, 4 + len(periodos)):
                ws.cell(fila, j).border = BORDE
        fila += 3
    ws.freeze_panes = "C4"
    _impresion(ws, 4)

    # ---------- Plan de mejora (ISO 9001: 9.1.3 y 10.2)
    ws = wb.create_sheet("Plan de mejora")
    cols = ["Código", "Indicador", "Área", "Período", "Resultado", "Meta", "Semáforo", "Estado",
            "Análisis de resultados", "Correcciones / plan de mejora", "Responsable de la carga"]
    _titulo(ws, "Desviaciones y acciones correctivas (ISO 9001:2015, 9.1.3 y 10.2)", sub, len(cols))
    _encabezado(ws, 4, cols, [10, 36, 20, 16, 13, 12, 12, 13, 50, 50, 24])
    desv = [m for m in ordenadas if m.semaforo != "verde"]
    for k, m in enumerate(sorted(desv, key=lambda m: (m.semaforo != "rojo", m.indicador.codigo, m.periodo)), 5):
        i = m.indicador
        dec = m.decimales if m.decimales is not None else 2
        fila = [i.codigo, i.nombre, i.area.nombre, calc.etiqueta(i.frecuencia, m.periodo), m.resultado, m.meta, None,
                ESTADO[m.estado], m.analisis, m.plan_mejora, m.cargado_por.nombre]
        for j, v in enumerate(fila, 1):
            c = ws.cell(k, j, v)
            c.font = Font(name=FUENTE, size=10)
            c.border = BORDE
            c.alignment = Alignment(vertical="top", wrap_text=j in (2, 9, 10))
        for j in (5, 6):
            ws.cell(k, j).number_format = _fmt_num(i.unidad, dec)
        _celda_sem(ws.cell(k, 7), m.semaforo)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(5, 4 + len(desv))}"
    _impresion(ws, 4)

    # ---------- Fichas técnicas
    ws = wb.create_sheet("Fichas técnicas")
    cols = ["Código", "Indicador", "Área", "Seguimiento / proceso", "Líder de proceso", "Resp. medición",
            "Objetivo", "Fuente de datos", "Unidad", "Frecuencia", "Sentido",
            "Meta (verde)", "Límite crítico", "Variables", "Constantes", "Fórmula", "Fórmula (en palabras)",
            "Método / notas", "Estado"]
    _titulo(ws, "Fichas técnicas de los indicadores", sub, len(cols))
    _encabezado(ws, 4, cols, [10, 36, 20, 34, 22, 22, 40, 28, 10, 13, 15, 12, 13, 40, 28, 34, 46, 40, 11])
    for k, i in enumerate(sorted(indicadores, key=lambda x: x.codigo), 5):
        consts = i.constantes or []
        fila = [i.codigo, i.nombre, i.area.nombre, i.ref_nombre or "", i.lider or "", i.resp_medicion or "",
                i.objetivo, i.fuente or "", i.unidad, i.frecuencia.capitalize(),
                "Mayor es mejor" if i.sentido == "mayor" else "Menor es mejor", i.meta, i.limite,
                "\n".join(f"• {v['etiqueta']} ({v['clave']})" for v in i.variables),
                "\n".join(f"• {c['etiqueta']} = {c['valor']}" for c in consts), i.formula,
                _legible(i.formula, [*i.variables, *consts]), i.notas or "", "Activo" if i.activo else "Archivado"]
        for j, v in enumerate(fila, 1):
            c = ws.cell(k, j, v)
            c.font = Font(name="Consolas" if j == 16 else FUENTE, size=10)
            c.border = BORDE
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if k % 2:
                c.fill = _fill(ZEBRA)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(5, 4 + len(indicadores))}"
    _impresion(ws, 4)
    return _guardar(wb)


# ------------------------------------------------------------------ bitácora
def reporte_bitacora(registros, codigos: dict, generado_por: str, verificacion: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bitácora"
    ahora = _local(datetime.utcnow())
    estado = "ÍNTEGRA ✔" if verificacion["ok"] else f"ALTERADA ✖ (registro #{verificacion['primer_error_id']})"
    cols = ["#", "Fecha y hora", "Usuario", "Rol", "Acción", "Entidad", "Indicador", "Detalle", "Hash SHA-256"]
    _titulo(ws, "Bitácora de trazabilidad del sistema",
            f"Generado el {ahora:%d/%m/%Y %H:%M} por {generado_por} · Cadena de hashes: {estado} · "
            f"{verificacion['total']} registros", len(cols))
    _encabezado(ws, 4, cols, [8, 17, 28, 11, 22, 12, 11, 90, 24])
    for k, b in enumerate(registros, 5):
        det = "; ".join(f"{a}: {v}" for a, v in (b.detalle or {}).items())
        fila = [b.id, _local(b.fecha), b.usuario_nombre, b.rol, b.accion, b.entidad,
                codigos.get(b.indicador_id, ""), det, b.hash[:16] + "…"]
        for j, v in enumerate(fila, 1):
            c = ws.cell(k, j, v)
            c.font = Font(name="Consolas" if j == 9 else FUENTE, size=10, color=GRIS if j == 9 else "000000")
            c.border = BORDE
            c.alignment = Alignment(vertical="top", wrap_text=(j == 8))
            if k % 2:
                c.fill = _fill(ZEBRA)
        ws.cell(k, 2).number_format = "dd/mm/yyyy hh:mm:ss"
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(5, 4 + len(registros))}"
    _impresion(ws, 4)
    return _guardar(wb)


# ------------------------------------------------------------------ ficha técnica (formato SGC)
def meta_texto(i) -> str:
    """Meta como se escribía en la ficha en papel: «96 Mín.» o «3 Máx.»."""
    return f"{i.meta:g} {'Mín.' if i.sentido == 'mayor' else 'Máx.'}"


def _formula_doc(i) -> str:
    """Fórmula de cálculo con la leyenda de símbolos (SIGLA = …; X = significado)."""
    f = i.formula
    for v in sorted([*i.variables, *(i.constantes or [])], key=lambda v: -len(v["clave"])):
        f = re.sub(r"\b" + re.escape(v["clave"]) + r"\b", v["clave"].upper(), f)
    lineas = [f"{i.sigla or 'IND'} = {f}"]
    lineas += [f"{v['clave'].upper()} = {v['etiqueta']}" for v in i.variables]
    lineas += [f"{c['clave'].upper()} = {c['etiqueta']} ({c['valor']:g})" for c in (i.constantes or [])]
    return "\n".join(lineas)


def _franjas(i):
    """(mín, máx) de cada franja del semáforo, como en la tabla «Min./Max.» de la ficha."""
    mayor = i.sentido == "mayor"
    verde = (i.meta, i.meta_tope) if mayor else (i.meta_tope, i.meta)
    amar = (i.limite, i.limite_tope) if mayor else (i.limite_tope, i.limite)
    rojo = (f"< {i.limite:g}", f"> {i.limite_tope:g}" if i.limite_tope is not None else "") if mayor else \
           (f"< {i.limite_tope:g}" if i.limite_tope is not None else "", f"> {i.limite:g}")
    return verde, amar, rojo


def _banda(ws, fila, texto, ncols=13):
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ncols)
    c = ws.cell(fila, 1, texto)
    c.font = Font(name=FUENTE, bold=True, size=10, color="FFFFFF")
    c.fill = _fill(NAVY)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[fila].height = 20


def _par(ws, fila, c1, etiqueta, c2, c3, valor, alto=None, mono=False, wrap=True):
    """Celda de etiqueta (c1) + valor combinado (c2..c3)."""
    e = ws.cell(fila, c1, etiqueta)
    e.font = Font(name=FUENTE, bold=True, size=10, color=NAVY)
    e.fill = _fill("E2E8F0")
    e.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    e.border = BORDE
    if c3 > c2:
        ws.merge_cells(start_row=fila, start_column=c2, end_row=fila, end_column=c3)
    v = ws.cell(fila, c2, valor)
    v.font = Font(name="Consolas" if mono else FUENTE, size=10)
    v.alignment = Alignment(vertical="top" if mono else "center", wrap_text=wrap, indent=1)
    for c in range(c2, c3 + 1):
        ws.cell(fila, c).border = BORDE
    if alto:
        ws.row_dimensions[fila].height = alto
    return v


def _lineas(texto: str, ancho: int) -> int:
    return sum(max(1, -(-len(t) // ancho)) for t in (texto or "").split("\n"))


def fichas_xlsx(fichas, params: dict, generado_por: str) -> bytes:
    """fichas: [(Indicador, ejercicio, periodos[{periodo, etiqueta, corta, medicion}])]. Una hoja por indicador."""
    wb = Workbook()
    wb.remove(wb.active)
    usados: set[str] = set()
    ahora = _local(datetime.utcnow())
    for ind, y, periodos in fichas:
        base = re.sub(r"[\\/*?:\[\]]", "", ind.sigla or ind.codigo)[:28] or ind.codigo
        nombre, n = base, 2
        while nombre.lower() in usados:
            nombre, n = f"{base[:26]}_{n}", n + 1
        usados.add(nombre.lower())
        ws = wb.create_sheet(nombre)
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 30
        for j in range(2, 14):
            ws.column_dimensions[get_column_letter(j)].width = 11.5
        tipo = {"proceso": "Desempeño de proceso", "objetivo": "Objetivo organizacional",
                "proveedor": "Desempeño de proveedor externo"}.get(ind.seguimiento or "proceso", "")
        etq1, etq2 = {"objetivo": ("OBJETIVO ORGANIZACIONAL", "OBJETIVO ESPECÍFICO"),
                      "proveedor": ("PROCESO", "BIEN o SERVICIO")}.get(ind.seguimiento or "proceso",
                                                                       ("PROCESO", "OBJETIVO DEL PROCESO"))
        dec = ind.decimales if ind.decimales is not None else 2
        fmt = _fmt_num(ind.unidad, dec)

        # --- encabezado
        ws.merge_cells("A1:C2")
        c = ws["A1"]
        c.value = params.get("empresa", "")
        c.font = Font(name=FUENTE, bold=True, size=14, color="FFFFFF")
        c.fill = _fill(NAVY)
        c.alignment = Alignment(vertical="center", indent=1, wrap_text=True)
        ws.merge_cells("D1:J2")
        c = ws["D1"]
        c.value = "FICHA TÉCNICA DE INDICADOR DE GESTIÓN"
        c.font = Font(name=FUENTE, bold=True, size=14, color="FFFFFF")
        c.fill = _fill(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells("K1:M1")
        ws["K1"].value = "Seguimiento de:"
        ws.merge_cells("K2:M2")
        ws["K2"].value = tipo
        for ref, bold in (("K1", False), ("K2", True)):
            ws[ref].font = Font(name=FUENTE, size=10, bold=bold, color="FFFFFF")
            ws[ref].fill = _fill(BRAND)
            ws[ref].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = ws.row_dimensions[2].height = 20
        ws.merge_cells("A3:M3")
        ws["A3"].value = params.get("gerencia_emisora", "")
        ws["A3"].font = Font(name=FUENTE, italic=True, size=10, color=GRIS)
        ws["A3"].alignment = Alignment(indent=1)

        # --- proceso / objetivo
        _banda(ws, 4, etq1)
        ws.merge_cells("A5:M5")
        ws["A5"].value = ind.ref_nombre or ""
        ws["A5"].font = Font(name=FUENTE, bold=True, size=11)
        ws["A5"].alignment = Alignment(indent=1, vertical="center")
        _banda(ws, 6, etq2)
        ws.merge_cells("A7:M7")
        ws["A7"].value = ind.ref_detalle or ind.objetivo or ""
        ws["A7"].font = Font(name=FUENTE, size=10)
        ws["A7"].alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.row_dimensions[7].height = max(18, 14 * _lineas(ws["A7"].value, 150))

        # --- definición + semáforo
        _banda(ws, 8, "DEFINICIÓN DEL INDICADOR E INFORMACIÓN PARA LA MEDICIÓN")
        nombre_ind = ind.nombre + (f" ({ind.sigla})" if ind.sigla and f"({ind.sigla})" not in ind.nombre else "")
        _par(ws, 9, 1, "Nombre del indicador", 2, 7, nombre_ind, 30)
        _par(ws, 10, 1, "Unidad de medida", 2, 3, ind.unidad)
        _par(ws, 10, 4, "Meta", 5, 7, meta_texto(ind))
        _par(ws, 11, 1, "Frecuencia de medición", 2, 7, ind.frecuencia.capitalize())
        doc = _formula_doc(ind)
        _par(ws, 12, 1, "Fórmula de cálculo", 2, 7, doc, mono=True)
        ws.merge_cells(start_row=12, start_column=2, end_row=14, end_column=7)
        for r in (12, 13, 14):
            ws.row_dimensions[r].height = max(20, 14 * _lineas(doc, 60) / 3 + 4)
        ws.merge_cells("A12:A14")
        _par(ws, 15, 1, "Fuente de la información", 2, 7, ind.fuente or "", 22)
        # semáforo
        ws.merge_cells("I9:M9")
        c = ws["I9"]
        c.value = "SEMÁFORO"
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF", size=10)
        c.fill = _fill(BRAND)
        c.alignment = Alignment(horizontal="center", vertical="center")
        for col, t in ((9, ""), (10, "Mín."), (12, "Máx.")):
            if col == 9:
                continue
            ws.merge_cells(start_row=10, start_column=col, end_row=10, end_column=col + 1)
            h = ws.cell(10, col, t)
            h.font = Font(name=FUENTE, bold=True, size=9, color=GRIS)
            h.alignment = Alignment(horizontal="center")
        verde, amar, rojo = _franjas(ind)
        for r, (nom, (mn, mx), clave) in enumerate((("Verde", verde, "verde"), ("Amarillo", amar, "amarillo"),
                                                    ("Rojo", rojo, "rojo")), 11):
            bg, fg, _ = SEM[clave]
            e = ws.cell(r, 9, nom)
            e.fill, e.font = _fill(bg), Font(name=FUENTE, bold=True, size=10, color=fg)
            e.alignment = Alignment(horizontal="center", vertical="center")
            e.border = BORDE
            for col, val in ((10, mn), (12, mx)):
                ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)
                v = ws.cell(r, col, val if val is not None else "")
                v.number_format = fmt
                v.alignment = Alignment(horizontal="center", vertical="center")
                v.font = Font(name=FUENTE, size=10)
                for cc in (col, col + 1):
                    ws.cell(r, cc).border = BORDE
        _par(ws, 15, 9, "Ejercicio fiscal", 10, 13, calc.etiqueta_ejercicio(y))

        # --- responsabilidades
        _banda(ws, 16, "RESPONSABILIDADES")
        _par(ws, 17, 1, "Gerencia", 2, 4, ind.gerencia or "", 22)
        _par(ws, 17, 5, "Departamento", 6, 8, ind.area.nombre)
        _par(ws, 17, 9, "Líder de proceso", 10, 13, ind.lider or "")
        _par(ws, 18, 1, "Responsable de medición", 2, 4, ind.resp_medicion or "", 22)
        _par(ws, 18, 5, "Responsable de registro", 6, 8, ind.resp_registro or "")
        _par(ws, 18, 9, "Responsable de análisis", 10, 13, ind.resp_analisis or "")

        # --- resultados de las determinaciones
        _banda(ws, 19, "RESULTADOS DE LAS DETERMINACIONES")
        hdr = ws.cell(20, 1, "PERÍODOS")
        hdr.font = Font(name=FUENTE, bold=True, size=10, color=NAVY)
        hdr.fill = _fill("E2E8F0")
        hdr.border = BORDE
        for j in range(12):
            c = ws.cell(20, 2 + j, periodos[j]["corta"] if j < len(periodos) else "No aplica")
            c.font = Font(name=FUENTE, bold=True, size=9, color="FFFFFF" if j < len(periodos) else GRIS)
            c.fill = _fill(BRAND if j < len(periodos) else "E2E8F0")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BORDE
        fila = 21
        for v in ind.variables:
            e = ws.cell(fila, 1, v["etiqueta"])
            e.font = Font(name=FUENTE, size=10)
            e.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
            e.border = BORDE
            ws.row_dimensions[fila].height = max(18, 14 * _lineas(v["etiqueta"], 30))
            for j in range(12):
                c = ws.cell(fila, 2 + j)
                c.border = BORDE
                c.alignment = Alignment(horizontal="center", vertical="center")
                if j < len(periodos) and periodos[j]["medicion"] and periodos[j]["medicion"].estado != "observado":
                    c.value = periodos[j]["medicion"].valores.get(v["clave"])
                    c.number_format = "#,##0.##"
                elif j >= len(periodos):
                    c.fill = _fill("F1F5F9")
            fila += 1

        # --- valores del indicador
        fila += 1
        _banda(ws, fila, "VALORES DEL INDICADOR")
        fila += 1
        for j, t in enumerate(["Período", "Indicador", "Mín.", "Máx.", "Semáforo"], 1):
            c = ws.cell(fila, j, t)
            c.font = Font(name=FUENTE, bold=True, size=10, color="FFFFFF")
            c.fill = _fill(NAVY)
            c.alignment = Alignment(horizontal="center")
            c.border = BORDE
        ini_tabla = fila + 1
        mayor = ind.sentido == "mayor"
        for per in periodos:
            fila += 1
            m = per["medicion"] if per["medicion"] and per["medicion"].estado != "observado" else None
            mn = ind.meta if mayor else ind.meta_tope
            mx = ind.meta_tope if mayor else ind.meta
            vals = [per["corta"], m.resultado if m else None, mn, mx]
            for j, v in enumerate(vals, 1):
                c = ws.cell(fila, j, v)
                c.border = BORDE
                c.font = Font(name=FUENTE, size=10, bold=(j == 2))
                c.alignment = Alignment(horizontal="center" if j != 1 else "left", indent=1 if j == 1 else 0)
                if j > 1:
                    c.number_format = fmt
            if m:
                _celda_sem(ws.cell(fila, 5), m.semaforo)
            ws.cell(fila, 5).border = BORDE
        fin_tabla = fila
        if periodos:
            ch = LineChart()
            ch.title = ind.nombre[:60]
            ch.height, ch.width = 7.5, 15.5
            ch.add_data(Reference(ws, min_col=2, max_col=4, min_row=ini_tabla - 1, max_row=fin_tabla), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=ini_tabla, max_row=fin_tabla))
            estilos = (("0284C7", None, True), ("16A34A", "dash", False), ("DC2626", "dash", False))
            for sr, (col, dash, marca) in zip(ch.series, estilos):
                sr.graphicalProperties.line.solidFill = col
                sr.graphicalProperties.line.width = 28000 if marca else 19000
                if dash:
                    sr.graphicalProperties.line.dashStyle = dash
                sr.smooth = False
                if marca:
                    sr.marker.symbol, sr.marker.size = "circle", 7
                    sr.marker.graphicalProperties.solidFill = col
                else:
                    sr.marker.symbol = "none"
            ch.x_axis.delete = ch.y_axis.delete = False
            ch.legend.position = "b"
            ws.add_chart(ch, f"G{ini_tabla - 1}")
        fila = max(fila, ini_tabla + 13)

        # --- análisis y evaluación
        fila += 1
        _banda(ws, fila, "ANÁLISIS Y EVALUACIÓN DE RESULTADOS")
        fila += 1
        ws.cell(fila, 1, "Período")
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=7)
        ws.cell(fila, 2, "Resultados del análisis")
        ws.merge_cells(start_row=fila, start_column=8, end_row=fila, end_column=13)
        ws.cell(fila, 8, "Correcciones, oportunidades o necesidades de mejora, propuestas de cambios y adecuación de recursos")
        for j in range(1, 14):
            c = ws.cell(fila, j)
            c.font = Font(name=FUENTE, bold=True, size=10, color="FFFFFF")
            c.fill = _fill(NAVY)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDE
        ws.row_dimensions[fila].height = 32
        for per in periodos:
            fila += 1
            m = per["medicion"] if per["medicion"] and per["medicion"].estado != "observado" else None
            ws.cell(fila, 1, per["etiqueta"]).font = Font(name=FUENTE, bold=True, size=10)
            ws.cell(fila, 1).alignment = Alignment(vertical="top", indent=1, wrap_text=True)
            ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=7)
            ws.merge_cells(start_row=fila, start_column=8, end_row=fila, end_column=13)
            ws.cell(fila, 2, m.analisis if m else "")
            ws.cell(fila, 8, m.plan_mejora if m else "")
            for j in range(1, 14):
                c = ws.cell(fila, j)
                c.border = BORDE
                if j in (2, 8):
                    c.font = Font(name=FUENTE, size=10)
                    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
            ws.row_dimensions[fila].height = max(20, 14 * max(_lineas(m.analisis if m else "", 66),
                                                              _lineas(m.plan_mejora if m else "", 66)))

        # --- notas y pie
        fila += 2
        _banda(ws, fila, "NOTAS")
        fila += 1
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=13)
        notas = ind.notas or ""
        ws.cell(fila, 1, notas)
        ws.cell(fila, 1).alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.cell(fila, 1).font = Font(name=FUENTE, size=10)
        ws.row_dimensions[fila].height = max(30, 14 * _lineas(notas, 150))
        fila += 2
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=13)
        ws.cell(fila, 1, f"{params.get('codigo_formato', '')} · Generado el {ahora:%d/%m/%Y %H:%M} por {generado_por} · SIGT")
        ws.cell(fila, 1).font = Font(name=FUENTE, size=8, color=GRIS)
        ws.cell(fila, 1).alignment = Alignment(horizontal="right")
        ws.print_area = f"A1:M{fila}"
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.oddFooter.right.text = f"{params.get('codigo_formato', '')}"
    return _guardar(wb)
